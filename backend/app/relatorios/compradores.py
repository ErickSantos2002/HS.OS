"""
Relatório: quem comprou um produto — DataCoreHS (ERP Tiny)

⚠️ **Existe porque o modelo não consegue escrever a lista.** Em 30/09/2026 o
CEO pediu "todas as empresas, inclusive as que compraram 1 unidade" de Phoebus.
A Iris apurou certo (137 empresas, 757 unidades), mas ao escrever a tabela bateu
no teto de saída do modelo (8.192 tokens) e o turno inteiro se perdeu — três
vezes seguidas, com a Nina culpando "contexto estourado". Lista longa escrita por
LLM também é cara, lenta e sujeita a erro de digitação. Aqui a lista sai do banco
direto para a planilha; o agente só recebe o resumo.

⚠️ **A régua de venda é a do DataCoreHS: `gold.fato_vendas`.** Não é
reimplementada aqui — o fato já sai filtrado por `analytics/models/silver/vendas.sql`
(CFOP do ITEM em 5102/6102/5108/6108/7102, Emitida DANFE, sem marcador que
exclua segundo `silver.stg_marcadores`, fora a curadoria). A skill `faturamento`
documenta a mesma coisa; as âncoras dos testes são as dela.

Até 01/10/2026 a régua era montada aqui, lendo `CFOP_VALIDOS` e
`MARCADORES_INVALIDOS` de `tiny.configuracoes` — CFOP tirado da natureza da
nota e marcador por `LIKE` contra a lista. As duas réguas foram rodadas lado a
lado antes da troca, e o Phoebus bateu nas três âncoras (2025 = 217, 2024 =
194, 2020–2026 = 137 empresas / 757) **por outro caminho**: a régua antiga
contava as vendas 005725 e 004743 e abatia as devoluções 005745 e 004762; o dbt
tira as duas vendas pelo marcador ("NOTA REJEITADA", "REJEITADA SEFAZ" — o
radical `rejeit`, que a lista não tinha) e não sobra o que abater. Na base
inteira a antiga contava 27 notas que não são venda e deixava de fora 13 que
são (exportação 7102 e natureza vazia).

O que continua sendo nosso:

  1. Item pelo nome EXATO — `ILIKE` traz fonte, placa e impressora junto.
  2. Devolução (CFOP 1202/2202) só é abatida quando a venda que ela cita nas
     `observacoes` ESTÁ em `gold.fato_vendas`. Se a origem já saiu por
     marcador, abater de novo conta a mesma devolução duas vezes — o erro de
     30/09, em que "bruto cheio − devoluções" deu 230 Phoebus em 2025 no lugar
     de 217. ⚠️ O DataCoreHS não abate devolução nenhuma (decisão D8: a venda
     devolvida sai pelo marcador). Em 01/10/2026, 9 das 21 devoluções da base
     citavam venda ainda no fato — sem marcador —, e nelas o relatório e o
     painel divergem de propósito.

Valor é o do ITEM (`valor_total_item`), não o da nota: a nota tem frete e
outros produtos. ⚠️ O fato é recalculado às 05:00: nota emitida hoje entra
amanhã.
"""
import html as _html
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MODOS = ("raiz_cnpj", "nome")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITULO_FONT = Font(bold=True, size=14)

# A nota de devolução cita a de origem em texto livre: "faturado na DANFE de
# venda 006031", "faturado na NF de venda 4743", "nota de Remessa de Saída DANFE
# 5542". O número vem depois de DANFE/NF, com até ~25 caracteres no meio.
_REF = re.compile(r"(?:DANFE|NF)[^0-9]{0,25}(\d{3,6})", re.IGNORECASE)

# Sufixos societários que fazem a mesma empresa parecer duas quando se agrupa
# por nome ("NACIONAL GAS BUTANO DISTRIBUIDORA LTDA" × "... LTDA.").
_SUFIXOS = {"SA", "S", "A", "LTDA", "ME", "EPP", "EIRELI", "CIA"}


def referencia_da_devolucao(observacoes: str | None) -> str | None:
    """O número (com 6 dígitos) da nota que a devolução diz estar devolvendo."""
    m = _REF.search(observacoes or "")
    return m.group(1).zfill(6) if m else None


def normalizar_nome(nome: str | None) -> str:
    """Maiúsculas, sem acento, sem pontuação e sem sufixo societário."""
    s = unicodedata.normalize("NFKD", nome or "").encode("ascii", "ignore").decode()
    palavras = re.sub(r"[^A-Za-z0-9 ]", " ", s).upper().split()
    return " ".join(p for p in palavras if p not in _SUFIXOS)


def chave_de_grupo(cpf_cnpj: str | None, nome: str | None, modo: str) -> str:
    """A que "empresa" uma nota pertence.

    `raiz_cnpj`: os 8 primeiros dígitos juntam filiais (a Nacional Gás tem 25
    CNPJs). CPF fica inteiro — pessoa não tem filial. Sem documento, cai no nome.
    `nome`: só o nome normalizado. `consolidar` NÃO usa este modo sozinho — ver
    o comentário de lá.
    """
    if modo == "nome":
        return "nome:" + normalizar_nome(nome)
    dig = re.sub(r"\D", "", cpf_cnpj or "")
    if len(dig) == 14:
        return "cnpj:" + dig[:8]
    if len(dig) == 11:
        return "cpf:" + dig
    return "nome:" + normalizar_nome(nome)


def consolidar(vendas: list[dict], devolucoes: list[dict], modo: str) -> dict:
    """Junta vendas e devoluções abatidas por empresa. Função pura — é o que os
    testes exercitam.

    `vendas`: {numero, data, cpf_cnpj, nome, quantidade, valor}
    `devolucoes`: as mesmas chaves + `decisao` ("abater" ou o motivo de não
    abater). Só as `abater` mexem nos totais; as outras vão para a aba de
    conferência, com o porquê.
    """
    grupos: dict[str, dict] = {}

    # ⚠️ **"nome" começa pela raiz e só então junta por nome.** Agrupar direto
    # pelo nome SEPARA filiais que o cadastro grafou diferente — medido em
    # 30/09/2026: 153 "empresas" de Phoebus contra 137 por raiz, o contrário de
    # "juntar a mesma empresa". Aqui duas raízes viram uma só quando algum nome
    # normalizado delas coincide (união, então A~B e B~C juntam os três).
    pai: dict[str, str] = {}

    def raiz(k):
        while pai.setdefault(k, k) != k:
            k = pai[k]
        return k

    if modo == "nome":
        dono_do_nome: dict[str, str] = {}
        for linha in (*vendas, *devolucoes):
            k = raiz(chave_de_grupo(linha.get("cpf_cnpj"), linha.get("nome"), "raiz_cnpj"))
            n = normalizar_nome(linha.get("nome"))
            if n in dono_do_nome:
                a = raiz(dono_do_nome[n])
                if a != k:
                    pai[a] = k
            else:
                dono_do_nome[n] = k

    def grupo(linha):
        k = chave_de_grupo(linha.get("cpf_cnpj"), linha.get("nome"), "raiz_cnpj")
        if modo == "nome":
            k = raiz(k)
        g = grupos.get(k)
        if g is None:
            g = grupos[k] = {"chave": k, "nomes": defaultdict(int),
                             "documentos": set(), "notas": set(), "brutas": 0,
                             "valor_bruto": 0.0, "devolvidas": 0,
                             "valor_devolvido": 0.0, "primeira": None, "ultima": None}
        if linha.get("cpf_cnpj"):
            g["documentos"].add(linha["cpf_cnpj"])
        g["nomes"][linha.get("nome") or "(sem nome)"] += 1
        return g

    por_ano: dict[int, dict] = defaultdict(lambda: {"brutas": 0, "valor_bruto": 0.0,
                                                    "devolvidas": 0, "valor_devolvido": 0.0})
    for v in vendas:
        g = grupo(v)
        g["notas"].add(v["numero"])
        g["brutas"] += int(v["quantidade"])
        g["valor_bruto"] += float(v["valor"])
        d = v["data"]
        g["primeira"] = d if g["primeira"] is None or d < g["primeira"] else g["primeira"]
        g["ultima"] = d if g["ultima"] is None or d > g["ultima"] else g["ultima"]
        a = por_ano[d.year]
        a["brutas"] += int(v["quantidade"])
        a["valor_bruto"] += float(v["valor"])

    for dv in devolucoes:
        if dv.get("decisao") != "abater":
            continue
        g = grupo(dv)
        g["devolvidas"] += int(dv["quantidade"])
        g["valor_devolvido"] += float(dv["valor"])
        a = por_ano[dv["data"].year]
        a["devolvidas"] += int(dv["quantidade"])
        a["valor_devolvido"] += float(dv["valor"])

    linhas = []
    for g in grupos.values():
        # O nome que mais aparece no grupo, não o da primeira nota — senão a
        # empresa aparece com o nome de uma filial ("JSL S/A. - Parauapebas").
        g["nome"] = min(g.pop("nomes").items(), key=lambda kv: (-kv[1], len(kv[0])))[0]
        g["liquidas"] = g["brutas"] - g["devolvidas"]
        g["valor_liquido"] = round(g["valor_bruto"] - g["valor_devolvido"], 2)
        linhas.append(g)
    linhas.sort(key=lambda g: (-g["liquidas"], -g["valor_liquido"], g["nome"]))

    return {
        "linhas": linhas,
        "por_ano": dict(sorted(por_ano.items())),
        "empresas": len(linhas),
        "brutas": sum(g["brutas"] for g in linhas),
        "devolvidas": sum(g["devolvidas"] for g in linhas),
        "liquidas": sum(g["liquidas"] for g in linhas),
        "valor_liquido": round(sum(g["valor_liquido"] for g in linhas), 2),
    }


# ─────────────────────────────────────────────────────────────────────────────
# Banco
# ─────────────────────────────────────────────────────────────────────────────

# O fato tem o item (`sk_venda_item` = `tiny.itens_nota.id`) mas não o nome dele
# nem o cliente da nota; esses vêm do `tiny`, por id. Nenhum filtro de régua aqui.
_VENDAS = """
SELECT f.numero_nota AS numero, f.data_venda AS data, c.cpf_cnpj, c.nome,
       f.quantidade, f.valor_total_item AS valor
  FROM gold.fato_vendas f
  JOIN tiny.itens_nota i ON i.id = f.sk_venda_item
  JOIN tiny.notas_fiscais nf ON nf.id = f.id_nota
  LEFT JOIN tiny.clientes c ON c.id = nf.id_cliente
 WHERE upper(btrim(i.descricao)) = upper(btrim(%(produto)s))
   AND f.data_venda BETWEEN %(inicio)s AND %(fim)s
"""

# A origem pode ser de um ano anterior ao período pedido: ela é procurada no fato
# inteiro, sem filtro de data. O que decide o ano do abatimento é a data da
# DEVOLUÇÃO.
_DEVOLUCOES = """
SELECT nf.numero, nf.data_emissao AS data, c.cpf_cnpj, c.nome,
       i.quantidade, i.valor_total AS valor, nf.observacoes,
       ARRAY(SELECT DISTINCT lpad(numero_nota, 6, '0') FROM gold.fato_vendas) AS contadas
  FROM tiny.notas_fiscais nf
  JOIN tiny.itens_nota i ON i.id_nota = nf.id
  LEFT JOIN tiny.clientes c ON c.id = nf.id_cliente
 WHERE upper(btrim(i.descricao)) = upper(btrim(%(produto)s))
   AND lower(btrim(nf.descricao_situacao)) = 'emitida danfe'
   AND substring(nf.natureza_operacao FROM '\\d{4}') IN ('1202', '2202')
   AND nf.data_emissao BETWEEN %(inicio)s AND %(fim)s
"""


def _decidir(ref: str | None, contadas: set[str]) -> str:
    if ref is None:
        return "não cita a nota de origem — não abatida"
    if ref not in contadas:
        return f"origem {ref} fora de gold.fato_vendas — não abatida (contaria duas vezes)"
    return "abater"


def _fetch(cur, sql, params):
    cur.execute(sql, params)
    cols = [c.name for c in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


def apurar(dsn: str, produto: str, inicio: date, fim: date,
           abater_devolucoes: bool = True, modo: str = "raiz_cnpj") -> dict:
    """Busca no DataCoreHS e consolida. Síncrono (psycopg): rodar em thread."""
    p = {"produto": produto, "inicio": inicio, "fim": fim}
    with psycopg.connect(dsn, connect_timeout=15) as conn, conn.cursor() as cur:
        vendas = _fetch(cur, _VENDAS, p)
        devolucoes = _fetch(cur, _DEVOLUCOES, p) if abater_devolucoes else []
    contadas = set(devolucoes[0]["contadas"]) if devolucoes else set()
    for d in devolucoes:
        d["referencia"] = referencia_da_devolucao(d.pop("observacoes"))
        d.pop("contadas", None)
        d["decisao"] = _decidir(d["referencia"], contadas)
    r = consolidar(vendas, devolucoes, modo)
    r["devolucoes"] = devolucoes
    return r


# ─────────────────────────────────────────────────────────────────────────────
# Planilha
# ─────────────────────────────────────────────────────────────────────────────

def _cabecalho(ws, linha, titulos):
    for i, t in enumerate(titulos, 1):
        c = ws.cell(row=linha, column=i, value=t)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _larguras(ws, larguras):
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def gerar(dsn: str, produto: str, inicio: date, fim: date,
          abater_devolucoes: bool = True, modo: str = "raiz_cnpj") -> tuple[bytes, str, dict]:
    """Apura e devolve (bytes do .xlsx, nome do arquivo, resumo)."""
    r = apurar(dsn, produto, inicio, fim, abater_devolucoes, modo)
    agrup = "raiz de CNPJ (filiais juntas)" if modo == "raiz_cnpj" else \
            "nome normalizado (aproximação: junta raízes diferentes do mesmo nome)"

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = f"Compradores de {produto}"
    ws["A1"].font = TITULO_FONT
    info = [
        ("Período", f"{inicio:%d/%m/%Y} a {fim:%d/%m/%Y}"),
        ("Empresas", r["empresas"]),
        ("Unidades vendidas", r["brutas"]),
        ("Unidades devolvidas (abatidas)", r["devolvidas"]),
        ("Unidades líquidas", r["liquidas"]),
        ("Valor líquido (R$)", r["valor_liquido"]),
        ("Agrupamento", agrup),
        ("Régua", "a do DataCoreHS (gold.fato_vendas): CFOP de venda do item, inclusive "
                  "exportação; Emitida DANFE; sem marcador que exclua. Item pelo nome "
                  "exato; valor do item. Dados até a carga das 05:00."),
        ("Devolução", "abatida só quando a venda de origem está em gold.fato_vendas"
                      if abater_devolucoes else "não abatida (pedido assim)"),
        ("Fonte", f"DataCoreHS / ERP Tiny — gerado em {date.today():%d/%m/%Y}"),
    ]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    _larguras(ws, [32, 90])

    ws = wb.create_sheet("Compradores")
    _cabecalho(ws, 1, ["#", "Empresa", "Documento(s)", "Nº de CNPJs/CPFs", "Notas",
                       "Un. vendidas", "Un. devolvidas", "Un. líquidas",
                       "Valor líquido (R$)", "Primeira compra", "Última compra"])
    for n, g in enumerate(r["linhas"], 1):
        docs = sorted(g["documentos"])
        ws.append([n, g["nome"], ", ".join(docs[:5]) + (" …" if len(docs) > 5 else ""),
                   len(docs), len(g["notas"]), g["brutas"], g["devolvidas"],
                   g["liquidas"], g["valor_liquido"], g["primeira"], g["ultima"]])
    ws.append([])
    ws.append(["", "TOTAL", "", "", "", r["brutas"], r["devolvidas"], r["liquidas"],
               r["valor_liquido"]])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    _larguras(ws, [5, 48, 40, 10, 8, 10, 10, 10, 16, 13, 13])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Por ano")
    _cabecalho(ws, 1, ["Ano", "Un. vendidas", "Un. devolvidas", "Un. líquidas",
                       "Valor líquido (R$)"])
    for ano, a in r["por_ano"].items():
        ws.append([ano, a["brutas"], a["devolvidas"], a["brutas"] - a["devolvidas"],
                   round(a["valor_bruto"] - a["valor_devolvido"], 2)])
    _larguras(ws, [8, 12, 12, 12, 18])

    ws = wb.create_sheet("Devoluções")
    _cabecalho(ws, 1, ["Nota", "Data", "Cliente", "Unidades", "Valor (R$)",
                       "Nota de origem", "Decisão"])
    for d in r["devolucoes"]:
        ws.append([d["numero"], d["data"], d["nome"], int(d["quantidade"]),
                   float(d["valor"]), d["referencia"] or "—", d["decisao"]])
    _larguras(ws, [10, 12, 44, 10, 14, 14, 60])

    buf = io.BytesIO()
    wb.save(buf)
    slug = re.sub(r"[^a-z0-9]+", "-", normalizar_nome(produto).lower()).strip("-")
    nome = f"compradores_{slug}_{inicio:%Y%m%d}-{fim:%Y%m%d}.xlsx"
    resumo = {k: r[k] for k in ("empresas", "brutas", "devolvidas", "liquidas",
                                "valor_liquido")}
    resumo["top"] = [(g["nome"], g["liquidas"], g["valor_liquido"])
                     for g in r["linhas"][:10]]
    resumo["abatidas"] = [d["numero"] for d in r["devolucoes"] if d["decisao"] == "abater"]
    resumo["html"] = pagina_html(produto, inicio, fim, r, agrup)
    return buf.getvalue(), nome, resumo


def _brl(v: float) -> str:
    return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pagina_html(produto: str, inicio: date, fim: date, r: dict, agrup: str) -> str:
    """A mesma lista da planilha, em HTML, montada aqui e não pelo modelo.

    ⚠️ **Em 30/09/2026 a Iris publicou a lista escrevendo o HTML ela mesma**,
    com a planilha certa já pronta ao lado. As 137 linhas somavam 768 unidades
    (o bruto) sob um cabeçalho que dizia "757 líquidas", e o valor não batia nem
    com o bruto nem com o líquido. Lista copiada por LLM erra; por isso o link
    também sai daqui, dos mesmos números que a planilha e os totais.
    """
    e = _html.escape
    linhas = "".join(
        f"<tr><td>{n}</td><td>{e(g['nome'])}</td><td class=n>{len(g['documentos'])}</td>"
        f"<td class=n>{g['brutas']}</td><td class=n>{g['devolvidas']}</td>"
        f"<td class=n><b>{g['liquidas']}</b></td><td class=n>{_brl(g['valor_liquido'])}</td></tr>"
        for n, g in enumerate(r["linhas"], 1))
    return f"""<!doctype html><html lang=pt-BR><meta charset=utf-8>
<meta name=viewport content="width=device-width,initial-scale=1">
<title>Compradores de {e(produto)}</title>
<style>body{{font-family:system-ui,sans-serif;margin:24px;color:#111}}
table{{border-collapse:collapse;width:100%;font-size:14px}}
th,td{{border-bottom:1px solid #ddd;padding:6px 8px;text-align:left}}
th{{background:#1f4e78;color:#fff;position:sticky;top:0}}.n{{text-align:right}}
.k{{display:inline-block;margin:0 24px 12px 0}}.k b{{font-size:22px;display:block}}
small{{color:#666}}</style>
<h1>Compradores de {e(produto)}</h1>
<p>{inicio:%d/%m/%Y} a {fim:%d/%m/%Y} · agrupado por {e(agrup)}</p>
<div><span class=k><b>{r['empresas']}</b>empresas</span>
<span class=k><b>{r['liquidas']}</b>unidades líquidas</span>
<span class=k><b>{_brl(r['valor_liquido'])}</b>valor líquido</span>
<span class=k><b>{r['brutas']}</b>vendidas · {r['devolvidas']} devolvidas</span></div>
<table><tr><th>#</th><th>Empresa</th><th class=n>CNPJs</th><th class=n>Vendidas</th>
<th class=n>Devolvidas</th><th class=n>Líquidas</th><th class=n>Valor líquido</th></tr>
{linhas}
<tr><th></th><th>TOTAL</th><th></th><th class=n>{r['brutas']}</th><th class=n>{r['devolvidas']}</th>
<th class=n>{r['liquidas']}</th><th class=n>{_brl(r['valor_liquido'])}</th></tr></table>
<p><small>Fonte: DataCoreHS (gold.fato_vendas, carga das 05:00). Só venda (CFOP de
venda do item, NF-e emitida, sem marcador de cancelamento), item pelo nome exato,
devolução abatida só quando a venda de origem está no fato. Gerado em
{date.today():%d/%m/%Y}.</small></p></html>"""
