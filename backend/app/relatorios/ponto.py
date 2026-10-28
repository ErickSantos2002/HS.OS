# backend/app/relatorios/ponto.py
"""Relatório de ponto do Flash — régua v2, Grupo A (Pontualidade e Jornada).

⚠️ **A régua é da coordenadora de RH, não nossa.** Mandada por escrito em
27/10/2026; mudou na conversa com ela, muda aqui — e muda `VERSAO_REGUA`, senão
o histórico compara meses apurados com réguas diferentes como se fossem iguais.
Ver `docs/superpowers/specs/2026-10-28-regua-ponto-v2-design.md`.

⚠️ **Por que existe código e não o agente:** na homologação o `bruce` leu os
CSV com `ler_anexo` e contou de cabeça — 21 eventos aprovados e 2 cancelados,
quando o arquivo tem 20 e 3. Número sai daqui.

Critérios **gerenciais**: o Flash faz a apuração oficial e nenhum saldo é
recalculado aqui.

⚠️ `CID` e `Justificativa` de evento **não são lidos** — `identificar` descarta
as colunas na leitura, antes de qualquer outra coisa. Ela pediu explicitamente
para não reproduzir diagnóstico.

**v2.1 (28/10/2026) — Cálculo Diário do FLASH, quarto arquivo, opcional mas
recomendado.** Com ele o calendário de cada pessoa é o do FLASH (`Natureza
Dia`), as batidas são as `Marcações Válidas` (tratadas por ele), e o dia útil
sem marcação com evento que só o FLASH conhece sai do denominador. O que
custou: o export de Eventos de setembro não trazia cinco dias de férias que o
FLASH conhecia, e eles apareciam como "sem marcação". ⚠️ **O FLASH não trata
intervalo sozinho** — num dia de 2 batidas ele conta o dia inteiro —, então o
Cálculo não resolve intervalo pendente: dá a evidência. `CPF`, `PIS`, `CNPJ` e
`Admissão` são descartados na leitura, como o `CID`.

Módulo puro: bytes entram, `Relatorio` sai; o histórico é um `dict` montado
aqui e gravado no R2 por `app/routers/mcp_anexos.py`.
"""
from __future__ import annotations

import calendar
import csv
import hashlib
import io
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta

# .1 = v2.1: o Cálculo Diário muda o calendário e a exclusão de dias.
VERSAO_REGUA = "2026-10-28.1"
TETO_LINHAS = 5000


def _hm(texto: str) -> int:
    h, m = texto.split(":")[:2]
    return int(h) * 60 + int(m)


# ── a régua (texto dela, 27/10/2026) — por minuto, segundos descartados ────
ENTRADA_ANTECIPADA_ANTES = _hm("07:20")  # entrada <  07:20 → A1
ATRASO_DESDE = _hm("07:46")              # entrada >= 07:46 → A2
SAIDA_ANTECIPADA_ANTES = _hm("16:50")    # saída   <  16:50 → A3
APOS_JORNADA_DEPOIS = _hm("17:10")       # saída   >  17:10 → A4
INTERVALO_MIN = 60                       # < A5 · = regular · > A6

MEDIDAS = ("entrada", "saida", "intervalo")
# ⚠️ A numeração A1–A6 é NOSSA, na ordem da tabela dela; ela fala em "A1 a A6"
# sem listar. Sai sempre número + nome, para ela corrigir se for outra ordem.
INDICADORES = {
    "A1": ("Entrada antecipada", "entrada"),
    "A2": ("Atraso", "entrada"),
    "A3": ("Saída antecipada", "saida"),
    "A4": ("Permanência após jornada", "saida"),
    "A5": ("Intervalo reduzido", "intervalo"),
    "A6": ("Intervalo prolongado", "intervalo"),
}
# (abaixo, faixa neutra, acima) de cada medida — é o que fecha a conta.
FAIXAS = {"entrada": ("A1", "Entrada permitida", "A2"),
          "saida": ("A3", "Saída permitida", "A4"),
          "intervalo": ("A5", "Intervalo regular", "A6")}
NEUTRA = "neutra"
NOME_MEDIDA = {"entrada": "entrada", "saida": "saída", "intervalo": "intervalo"}

IDENTIFICADA, JUSTIFICADA = "identificada", "justificada"

# Caractere de controle num campo do Flash faria o openpyxl recusar a célula
# e derrubar a planilha inteira; não há uso legítimo para eles aqui.
_CONTROLE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

_COL_MARC = {"Nome", "Data da Marcação", "Hora da Marcação 1", "Tipo Marcação 1"}
_COL_BANCO = {"Nome", "Matrícula", "Saldo banco total"}
_COL_EVENTOS = {"EventoId", "Motivo", "Data Inicio", "Status"}
# Só estas colunas do Eventos sobrevivem à leitura. `CID`, `Justificativa` e
# `Observação da Aprovação` ficam de fora por construção — o que não é lido não
# vaza por descuido numa aba nova.
_EVENTOS_USADAS = ("Matrícula", "Nome", "Departamento", "EventoId", "Motivo", "Descrição",
                   "Data Inicio", "Data Fim", "Hora Inicio", "Hora Fim", "Qtde", "Status")
_COL_CALCULO = {"Marcações Válidas", "Natureza Dia", "Jornada Considerada"}
# Idem no Cálculo Diário: `CPF`, `PIS`, `CNPJ` e `Admissão` não passam daqui.
_CALCULO_USADAS = ("Matrícula", "Nome", "Data", "Natureza Dia", "Nome Escala", "Marcações Válidas",
                   "Jornada Considerada", "Horas Trabalhadas", "Eventos")
TIPOS = {
    "marcacoes": "marcações (MarcacoesOriginaisJustificadas)",
    "eventos": "eventos (Eventos)",
    "banco": "banco de horas (BancoHoras)",
    "calculo": "cálculo diário (CalculoDiario)",
}


class ErroPonto(ValueError):
    """Problema nos arquivos, com mensagem que o agente repassa à pessoa."""


# ── calendário ──────────────────────────────────────────────────────────────

def pascoa(ano: int) -> date:
    """Domingo de Páscoa — algoritmo de Meeus/Jones/Butcher (gregoriano)."""
    a = ano % 19
    b, c = divmod(ano, 100)
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes, dia = divmod(h + l - 7 * m + 114, 31)
    return date(ano, mes, dia + 1)


# A lista dela: todos valem para a HS (Recife/PE).
_FERIADOS_FIXOS = (
    (1, 1, "Confraternização Mundial", "Nacional"),
    (3, 6, "Data Magna", "Estadual"),
    (4, 21, "Tiradentes", "Nacional"),
    (5, 1, "Dia do Trabalhador", "Nacional"),
    (6, 24, "São João", "Estadual"),
    (9, 7, "Independência", "Nacional"),
    (10, 12, "Nossa Senhora Aparecida", "Nacional"),
    (10, 19, "Dia do Comerciário", "Municipal"),
    (11, 2, "Finados", "Nacional"),
    (11, 15, "Proclamação da República", "Nacional"),
    (11, 20, "Consciência Negra", "Nacional"),
    (12, 8, "Imaculada Conceição", "Municipal"),
    (12, 25, "Natal", "Nacional"),
)


def feriados(ano: int) -> dict[date, tuple[str, str]]:
    """`{data: (nome, tipo)}`. Carnaval e Paixão saem da Páscoa — fixá-los em
    2026 quebraria a régua em janeiro."""
    p = pascoa(ano)
    saida = {date(ano, m, d): (nome, tipo) for m, d, nome, tipo in _FERIADOS_FIXOS}
    saida[p - timedelta(days=47)] = ("Carnaval", "Nacional")
    saida[p - timedelta(days=2)] = ("Paixão de Cristo", "Nacional")
    return saida


def dias_uteis(ano: int, mes: int) -> list[date]:
    fer = feriados(ano)
    d, saida = date(ano, mes, 1), []
    while d.month == mes:
        if d.weekday() < 5 and d not in fer:
            saida.append(d)
        d += timedelta(days=1)
    return saida


# ── estruturas ──────────────────────────────────────────────────────────────

@dataclass
class Evento:
    id: str
    chave: str
    matricula: str
    nome: str
    motivo: str
    codigo: str
    inicio: date
    fim: date
    hora_inicio: int | None
    hora_fim: int | None
    qtde: str
    status: str
    usado: bool = False
    por_que: str = ""
    justificou: list[str] = field(default_factory=list)

    @property
    def aprovado(self) -> bool:
        return self.status.strip().lower() == "aprovado"

    @property
    def dia_inteiro(self) -> bool:
        return self.hora_inicio is None and self.hora_fim is None

    @property
    def parcial(self) -> bool:
        return self.hora_inicio is not None and self.hora_fim is not None

    def evidencia(self) -> str:
        return f"{self.id} · {self.motivo}" + (f" ({self.codigo})" if self.codigo else "")


@dataclass
class Ocorrencia:
    chave: str
    matricula: str
    nome: str
    data: date
    indicador: str        # "A1".."A6"
    minutos: int          # horário da batida, ou duração do intervalo
    situacao: str         # identificada | justificada
    evento: str           # evidência do evento que justifica ("" se nenhum)
    batidas: str          # "07:52 · 12:00 · 13:00 · 17:05"
    ajuste: str           # "batida ajustada: <motivo>" ("" se original)
    flash_jornada: str = ""   # Cálculo Diário, quando veio
    flash_eventos: str = ""


@dataclass
class Pendencia:
    chave: str
    matricula: str
    nome: str
    data: date
    medida: str           # dia | saída | intervalo | fora da escala | aviso
    motivo: str
    # Evidência do Cálculo Diário daquele dia ("" sem o arquivo ou sem a linha).
    flash_jornada: str = ""
    flash_horas: str = ""
    flash_marcacoes: str = ""
    flash_eventos: str = ""


@dataclass
class CalcDia:
    """Uma linha do Cálculo Diário: o que o FLASH diz daquele pessoa-dia."""
    natureza: str         # Trabalho | Folga | Feriado
    marcacoes: list[int]  # Marcações Válidas, em minutos e em ordem
    marcacoes_txt: str    # como veio ("07:40 - 12:28 | 13:08 - 17:00", "-")
    jornada: str          # Jornada Considerada
    horas: str            # Horas Trabalhadas
    eventos: str          # nomes separados por vírgula, "-" quando nenhum

    @property
    def trabalho(self) -> bool:
        return self.natureza.strip().lower() == "trabalho"

    @property
    def tem_evento(self) -> bool:
        return self.eventos.strip() not in ("", "-")


@dataclass
class Aviso:
    """Diferença entre o Cálculo Diário e o resto — não é ocorrência nem pendência."""
    tipo: str             # calendário | marcações | pessoa-dia
    chave: str
    matricula: str
    nome: str
    data: date
    texto: str


@dataclass
class DiaAvaliado:
    """Um dia elegível de uma pessoa: a faixa de cada medida, ou `None` se pendente."""
    chave: str
    data: date
    faixa: dict           # medida -> "A1" | "neutra" | "A2" | None


@dataclass
class Pessoa:
    chave: str
    matricula: str
    nome: str
    setor: str
    dias_uteis: int = 0
    dias_fora: int = 0
    dias_fora_flash: int = 0     # dos `dias_fora`, os que só o Cálculo Diário conhecia
    dias_elegiveis: int = 0
    dias_pendentes: int = 0
    identificadas: dict = field(default_factory=lambda: dict.fromkeys(INDICADORES, 0))
    justificadas: dict = field(default_factory=lambda: dict.fromkeys(INDICADORES, 0))
    avaliados: dict = field(default_factory=lambda: dict.fromkeys(MEDIDAS, 0))
    pendentes: dict = field(default_factory=lambda: dict.fromkeys(MEDIDAS, 0))
    neutras: dict = field(default_factory=lambda: dict.fromkeys(MEDIDAS, 0))
    saldo_min: int | None = None
    saldo_txt: str = ""


@dataclass
class Conferencia:
    verificacao: str
    esperado: int
    obtido: int

    @property
    def ok(self) -> bool:
        return self.esperado == self.obtido


@dataclass
class Comparacao:
    status: str           # comparado | outra_versao | sem_historico | indisponivel
    periodo: str          # "AAAA-MM" do mês anterior
    mensagem: str
    linhas: dict = field(default_factory=dict)   # "A1" -> {identificadas, denominador, pct, variacao_pp}


@dataclass
class Relatorio:
    mes: int
    ano: int
    arquivos: list[tuple[str, str]]          # (nome, sha256)
    uteis: list[date]
    pessoas: list[Pessoa]
    dias: list[DiaAvaliado]
    ocorrencias: list[Ocorrencia]
    pendencias: list[Pendencia]
    eventos: list[Evento]
    tem_banco: bool
    conferencias: list[Conferencia] = field(default_factory=list)
    comparacao: Comparacao | None = None
    tem_calculo: bool = False
    avisos: list[Aviso] = field(default_factory=list)
    # trabalho, marcacoes_comparadas, marcacoes_iguais, dias_fora_flash
    calculo: dict = field(default_factory=dict)


# ── leitura ─────────────────────────────────────────────────────────────────

def _decodificar(dados: bytes) -> str:
    # O Flash exporta UTF-8 com BOM; reaberto e salvo no Excel vira Latin-1.
    try:
        return dados.decode("utf-8-sig")
    except UnicodeDecodeError:
        return dados.decode("latin-1")


def identificar(nome: str, dados: bytes) -> tuple[str, list[dict]]:
    """`("marcacoes" | "eventos" | "banco" | "calculo", registros)` — pelo cabeçalho, não pelo nome."""
    leitor = csv.DictReader(io.StringIO(_decodificar(dados)), delimiter=";")
    try:
        colunas = {(c or "").strip() for c in (leitor.fieldnames or [])}
    except csv.Error as e:
        raise ErroPonto(f"{nome} não pôde ser lido como CSV: {e}.") from e
    if _COL_MARC <= colunas:
        tipo = "marcacoes"
    elif _COL_EVENTOS <= colunas:
        tipo = "eventos"
    elif _COL_CALCULO <= colunas:
        tipo = "calculo"
    elif _COL_BANCO <= colunas:
        tipo = "banco"
    else:
        raise ErroPonto(
            f"{nome} não parece um export do Flash: esperava o CSV de marcações "
            "(MarcacoesOriginaisJustificadas), o de eventos (Eventos), o de banco de "
            "horas (BancoHoras) ou o de cálculo diário (CalculoDiario).")
    usadas = {"eventos": _EVENTOS_USADAS, "calculo": _CALCULO_USADAS}.get(tipo)
    try:
        registros = []
        for r in leitor:
            if not any((v or "").strip() for v in r.values() if isinstance(v, str)):
                continue
            linha = {(k or "").strip(): _CONTROLE.sub("", v if isinstance(v, str) else "").strip()
                     for k, v in r.items()}
            if usadas:
                linha = {k: linha.get(k, "") for k in usadas}
            registros.append(linha)
    except csv.Error as e:
        raise ErroPonto(f"{nome} não pôde ser lido como CSV: {e}.") from e
    if len(registros) > TETO_LINHAS:
        raise ErroPonto(f"{nome} tem {len(registros)} linhas; o limite é {TETO_LINHAS}.")
    return tipo, registros


def _minutos(texto: str) -> int | None:
    m = re.match(r"^\s*(\d{1,2}):(\d{2})", texto or "")
    return int(m.group(1)) * 60 + int(m.group(2)) if m else None


def _data(texto: str) -> date | None:
    """`dd/mm/aaaa` do Flash; qualquer outra forma é `None` (quem chama nomeia a linha)."""
    m = re.match(r"^(\d{2})/(\d{2})/(\d{4})$", (texto or "").strip())
    if not m or int(m.group(3)) < 2000:
        return None
    try:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def _saldo(texto: str) -> tuple[int | None, str]:
    """`(-07:42)` → (-462, "-07:42"); `02:05` → (125, "02:05")."""
    t = (texto or "").strip()
    negativo = t.startswith("(") or t.startswith("-")
    m = re.search(r"(\d+):(\d{2})", t)
    if not m:
        return None, ""
    total = int(m.group(1)) * 60 + int(m.group(2))
    return (-total if negativo else total), f"{'-' if negativo else ''}{m.group(1).zfill(2)}:{m.group(2)}"


def _matricula(texto: str) -> str:
    t = (texto or "").strip().lstrip("0")
    return t if t.isdigit() else ""


def _sem_acento(t: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", (t or "").lower())
                   if unicodedata.category(c) != "Mn").strip()


def _chave(matricula: str, nome: str) -> str:
    """Matrícula sem zeros à esquerda; sem matrícula (`-`), nome sem acento e minúsculo."""
    return f"m:{matricula}" if matricula else f"n:{_sem_acento(nome)}"


def _fmt(minutos: int | None) -> str:
    if minutos is None:
        return ""
    sinal = "-" if minutos < 0 else ""
    return f"{sinal}{abs(minutos) // 60:02d}:{abs(minutos) % 60:02d}"


def _codigo(descricao: str) -> str:
    """O número entre parênteses na Descrição: `"... do dia (1004)"` → `"1004"`."""
    achados = re.findall(r"\((\d{3,5})\)", descricao or "")
    return achados[-1] if achados else ""


# ── a régua ─────────────────────────────────────────────────────────────────

def faixa(medida: str, minutos: int) -> str:
    """`"A1".."A6"` ou `NEUTRA`. Os limites estão nas constantes acima."""
    if medida == "entrada":
        return "A1" if minutos < ENTRADA_ANTECIPADA_ANTES else "A2" if minutos >= ATRASO_DESDE else NEUTRA
    if medida == "saida":
        return "A3" if minutos < SAIDA_ANTECIPADA_ANTES else "A4" if minutos > APOS_JORNADA_DEPOIS else NEUTRA
    return "A5" if minutos < INTERVALO_MIN else "A6" if minutos > INTERVALO_MIN else NEUTRA


def _sobreposicao(periodos: list[tuple[int, int]], ini: int, fim: int) -> int:
    """Minutos de `[ini, fim]` cobertos pela união dos períodos."""
    cortes = sorted((max(a, ini), min(b, fim)) for a, b in periodos if min(b, fim) > max(a, ini))
    total, ate = 0, ini
    for a, b in cortes:
        a = max(a, ate)
        if b > a:
            total += b - a
            ate = b
    return total


def justificativa(indicador: str, batidas: list[int], eventos: list[Evento]) -> list[Evento]:
    """Os eventos aprovados parciais do dia que justificam a ocorrência (vazio = identificada).

    A1, A4 e A5 não são justificáveis por evento de ausência: chegar cedo, sair
    tarde e voltar antes não são ausência.
    """
    if indicador == "A2":
        entrada = batidas[0]
        return [e for e in eventos if e.hora_inicio <= ATRASO_DESDE - 1 and e.hora_fim >= entrada]
    if indicador == "A3":
        saida = batidas[-1]
        return [e for e in eventos if e.hora_inicio <= saida and e.hora_fim >= SAIDA_ANTECIPADA_ANTES]
    if indicador == "A6":
        ini, fim = batidas[1], batidas[2]
        quem = [e for e in eventos if min(e.hora_fim, fim) > max(e.hora_inicio, ini)]
        cobre = _sobreposicao([(e.hora_inicio, e.hora_fim) for e in quem], ini, fim)
        return quem if cobre > 0 and (fim - ini) - cobre <= INTERVALO_MIN else []
    return []


def _ler_eventos(registros: list[dict], nome_arq: str) -> list[Evento]:
    eventos = []
    for n, l in enumerate(registros, start=2):
        inicio = _data(l.get("Data Inicio", ""))
        if inicio is None:
            raise ErroPonto(f"{nome_arq}: data inválida {l.get('Data Inicio', '')!r} na linha {n} "
                            "(esperado dd/mm/aaaa).")
        fim_txt = l.get("Data Fim", "")
        fim = _data(fim_txt) if fim_txt else inicio
        if fim is None:
            raise ErroPonto(f"{nome_arq}: data inválida {fim_txt!r} na linha {n} (esperado dd/mm/aaaa).")
        mat = _matricula(l.get("Matrícula", ""))
        eventos.append(Evento(
            id=l.get("EventoId", ""), chave=_chave(mat, l.get("Nome", "")), matricula=mat,
            nome=l.get("Nome", ""), motivo=l.get("Motivo", ""), codigo=_codigo(l.get("Descrição", "")),
            inicio=inicio, fim=max(fim, inicio),
            hora_inicio=_minutos(l.get("Hora Inicio", "")), hora_fim=_minutos(l.get("Hora Fim", "")),
            qtde=l.get("Qtde", ""), status=l.get("Status", "")))
    return eventos


def _ler_calculo(registros: list[dict], nome_arq: str, mes: int, ano: int,
                 casar) -> dict[tuple[str, date], CalcDia]:
    """`{(chave, dia): CalcDia}` das pessoas que têm marcação no mês.

    Quem só existe no Cálculo (o CEO, sem matrícula e sem marcação) fica de
    fora, como na v2. O arquivo tem que ser do mesmo mês das marcações.
    """
    datas = []
    for n, l in enumerate(registros, start=2):
        d = _data(l.get("Data", ""))
        if d is None:
            raise ErroPonto(f"{nome_arq}: data inválida {l.get('Data', '')!r} na linha {n} "
                            "(esperado dd/mm/aaaa).")
        datas.append(d)
    outros = sorted({(d.month, d.year) for d in datas} - {(mes, ano)})
    if outros:
        rotulos = ", ".join(f"{m:02d}/{a}" for m, a in outros)
        raise ErroPonto(f"O CSV de Cálculo Diário ({nome_arq}) precisa ser do mesmo mês das "
                        f"marcações ({mes:02d}/{ano}); veio também: {rotulos}.")
    calc: dict[tuple[str, date], CalcDia] = {}
    linha_de: dict[tuple[str, date], int] = {}
    for n, (l, d) in enumerate(zip(registros, datas), start=2):
        dono = casar(_matricula(l.get("Matrícula", "")), l.get("Nome", ""))
        if dono is None:
            continue
        if (dono, d) in calc:
            raise ErroPonto(f"{nome_arq}: a linha {n} repete pessoa e dia da linha "
                            f"{linha_de[(dono, d)]} ({d:%d/%m/%Y}).")
        txt = l.get("Marcações Válidas", "")
        linha_de[(dono, d)] = n
        calc[(dono, d)] = CalcDia(
            natureza=l.get("Natureza Dia", ""),
            marcacoes=[_hm(x) for x in re.findall(r"\b\d{1,2}:\d{2}\b", txt)],
            marcacoes_txt=txt, jornada=l.get("Jornada Considerada", ""),
            horas=l.get("Horas Trabalhadas", ""), eventos=l.get("Eventos", ""))
    return calc


def _pausas(batidas: list[int]) -> str:
    """Pares (entrada, saída): as pausas são da saída de um à entrada do seguinte."""
    pausas = [(batidas[i], batidas[i + 1]) for i in range(1, len(batidas) - 1, 2)]
    return (f"FLASH registrou {len(pausas)} pausas: "
            + ", ".join(f"{_fmt(a)}–{_fmt(b)} ({b - a} min)" for a, b in pausas))


def _evidencia_ajuste(linha: dict, posicoes: list[int]) -> str:
    motivos = [linha.get(f"Motivo Marcação {i}", "") or "sem motivo"
               for i in posicoes if linha.get(f"Tipo Marcação {i}") == "Justificada"]
    return "; ".join(f"batida ajustada: {m}" for m in motivos)


def gerar(arquivos: list[tuple[str, bytes]]) -> Relatorio:
    """`arquivos` = [(nome, bytes)]: marcações e eventos (obrigatórios), banco e
    cálculo diário (opcionais; o cálculo, recomendado)."""
    lidos: dict[str, list[dict]] = {}
    nomes: dict[str, str] = {}
    hashes = []
    for nome, dados in arquivos:
        tipo, registros = identificar(nome, dados)
        if tipo in lidos:
            raise ErroPonto(f"Vieram dois CSV de {TIPOS[tipo]}: {nomes[tipo]} e {nome}. "
                            "Mande um de cada.")
        lidos[tipo], nomes[tipo] = registros, nome
        hashes.append((nome, hashlib.sha256(dados).hexdigest()))
    faltam = [TIPOS[t] for t in ("marcacoes", "eventos") if t not in lidos]
    if faltam:
        achou = ", ".join(nomes.values()) or "nenhum"
        raise ErroPonto(
            f"Falta o CSV de {' e '.join(faltam)} — é obrigatório: sem os eventos não há como "
            f"saber quem estava de férias, atestado ou day-off, e o denominador sai errado. "
            f"Reconhecido: {achou}.")

    linhas = lidos["marcacoes"]
    datas = []
    for n, l in enumerate(linhas, start=2):   # linha 1 é o cabeçalho
        d = _data(l.get("Data da Marcação", ""))
        if d is None:
            raise ErroPonto(f"{nomes['marcacoes']}: data inválida "
                            f"{l.get('Data da Marcação', '')!r} na linha {n} (esperado dd/mm/aaaa).")
        datas.append(d)
    meses = sorted({(d.month, d.year) for d in datas})
    if len(meses) != 1:
        rotulos = ", ".join(f"{m:02d}/{a}" for m, a in meses) or "nenhum"
        raise ErroPonto(f"O CSV de marcações precisa ser de um mês só; veio: {rotulos}.")
    mes, ano = meses[0]
    uteis = dias_uteis(ano, mes)
    uteis_set = set(uteis)
    fer = feriados(ano)

    # Pessoas: quem tem marcação no mês. Casa por matrícula; sem ela, por nome.
    pessoas: dict[str, Pessoa] = {}
    por_dia: dict[tuple[str, date], tuple[int, dict]] = {}
    for n, (l, dia) in enumerate(zip(linhas, datas), start=2):
        mat = _matricula(l.get("Matrícula", ""))
        chave = _chave(mat, l.get("Nome", ""))
        pessoas.setdefault(chave, Pessoa(chave, mat, l.get("Nome", ""), l.get("Departamento", "")))
        if (chave, dia) in por_dia:
            raise ErroPonto(f"{nomes['marcacoes']}: a linha {n} repete pessoa e dia da linha "
                            f"{por_dia[(chave, dia)][0]} ({dia:%d/%m/%Y}).")
        por_dia[(chave, dia)] = (n, l)
    por_nome = {_sem_acento(p.nome): p.chave for p in pessoas.values()}

    def casar(mat: str, nome: str) -> str | None:
        if mat and f"m:{mat}" in pessoas:
            return f"m:{mat}"
        return por_nome.get(_sem_acento(nome))

    primeiro_dia = date(ano, mes, 1)
    ultimo_dia = date(ano, mes, calendar.monthrange(ano, mes)[1])
    dias_do_mes = list(_intervalo(primeiro_dia, ultimo_dia))

    # v2.1 — Cálculo Diário (opcional). Sem ele, `calc` fica vazio e tudo
    # abaixo cai na v2: calendário da lista, batidas do CSV de marcações.
    tem_calculo = "calculo" in lidos
    calc = _ler_calculo(lidos["calculo"], nomes["calculo"], mes, ano, casar) if tem_calculo else {}
    avisos: list[Aviso] = []

    def avisar(tipo: str, chave: str, d: date, texto: str) -> None:
        p = pessoas[chave]
        avisos.append(Aviso(tipo, chave, p.matricula, p.nome, d, texto))

    def nosso_dia(d: date) -> str:
        if d in uteis_set:
            return "dia útil"
        if d in fer:
            return f"feriado ({fer[d][0]})"
        return "sábado" if d.weekday() == 5 else "domingo"

    # 1. Calendário do FLASH vence; a lista de feriados é o recuo do dia sem linha.
    util_por: dict[str, set[date]] = {}
    for chave in pessoas:
        util_por[chave] = set()
        for d in dias_do_mes:
            c = calc.get((chave, d))
            util = c.trabalho if c else d in uteis_set
            if c and util != (d in uteis_set):
                avisar("calendário", chave, d, f"FLASH diz {c.natureza or 'vazio'} em {d:%d/%m}, "
                                               f"nossa régua dizia {nosso_dia(d)}; vale o FLASH")
            if util:
                util_por[chave].add(d)

    # 2 e 5. Marcações Válidas × CSV de marcações, e pessoa-dia de um × do outro.
    batidas_de: dict[tuple[str, date], list[int]] = {}
    comparadas = iguais = 0
    if tem_calculo:
        for chave in pessoas:
            for d in dias_do_mes:
                c = calc.get((chave, d))
                no_csv = _batidas(por_dia.get((chave, d), (0, {}))[1])
                if c is None:
                    if no_csv:
                        avisar("pessoa-dia", chave, d, "dia com marcação no CSV de marcações que não "
                                                       "está no Cálculo Diário; valem as do CSV")
                    continue
                batidas_de[(chave, d)] = c.marcacoes
                if no_csv:
                    comparadas += 1
                    iguais += no_csv == c.marcacoes
                elif c.marcacoes:
                    avisar("pessoa-dia", chave, d, "o Cálculo Diário tem marcação ("
                                                   f"{_texto_batidas(c.marcacoes)}) e o CSV de "
                                                   "marcações não; valem as do FLASH")
                    continue
                if no_csv != c.marcacoes:
                    avisar("marcações", chave, d,
                           f"Marcações Válidas do FLASH ({_texto_batidas(c.marcacoes) or 'nenhuma'}) "
                           f"diferem do CSV de marcações ({_texto_batidas(no_csv)}); valem as do FLASH")

    # Eventos: o que tira dia do denominador e o que justifica período.
    eventos = _ler_eventos(lidos["eventos"], nomes["eventos"])
    dias_fora: dict[tuple[str, date], Evento] = {}
    parciais: dict[tuple[str, date], list[Evento]] = defaultdict(list)
    com_evento: set[tuple[str, date]] = set()   # qualquer evento, de qualquer status
    for e in eventos:
        dono = casar(e.matricula, e.nome)
        if dono:
            e.chave = dono
        # Recorta ao mês ANTES de andar: as datas vêm do arquivo, e um evento de
        # 2000 a 9999 seriam milhões de dias por linha (revisão de 28/10/2026).
        no_mes = list(_intervalo(max(e.inicio, primeiro_dia), min(e.fim, ultimo_dia)))
        if dono:
            com_evento.update((dono, d) for d in no_mes)
        uteis_ev = [d for d in no_mes if d in util_por.get(dono, uteis_set)]
        if not e.aprovado:
            e.por_que = f"não: status {e.status or 'vazio'} — só Aprovado justifica"
        elif dono is None:
            e.por_que = "não: pessoa sem marcação no mês"
        elif not no_mes:
            e.por_que = "não: fora do mês apurado"
        elif not (e.dia_inteiro or e.parcial):
            e.por_que = "não: só uma das horas (início/fim) preenchida — conferir no Flash"
        elif not uteis_ev:
            e.por_que = "não: nenhum dia útil do mês no período"
        elif e.dia_inteiro:
            e.usado = True
            for d in uteis_ev:
                dias_fora.setdefault((dono, d), e)
            e.por_que = f"sim: {len(uteis_ev)} dia(s) útil(eis) fora do denominador"
        else:
            e.usado = True
            for d in uteis_ev:
                parciais[(dono, d)].append(e)
            e.por_que = "sim: período justificado"

    dias: list[DiaAvaliado] = []
    ocorrencias: list[Ocorrencia] = []
    pendencias: list[Pendencia] = []

    def pendencia(p: Pessoa, d: date, medida: str, motivo: str) -> None:
        # 4. Toda pendência de um dia que está no Cálculo leva o que o FLASH diz.
        c = calc.get((p.chave, d))
        ev_flash = (c.jornada, c.horas, c.marcacoes_txt, c.eventos) if c else ("",) * 4
        pendencias.append(Pendencia(p.chave, p.matricula, p.nome, d, medida, motivo, *ev_flash))

    # Marcação fora da escala: sábado, domingo ou feriado (ou folga no FLASH)
    # não entra em indicador.
    for (chave, d), (_, l) in sorted(por_dia.items(), key=lambda x: (x[0][0], x[0][1])):
        if d in util_por[chave]:
            continue
        p = pessoas[chave]
        c = calc.get((chave, d))
        quando = nosso_dia(d) if d not in uteis_set else f"{c.natureza or 'dia não útil'} no FLASH"
        if _batidas(l):
            pendencia(p, d, "fora da escala", f"marcação fora da escala: {quando}")

    for p in pessoas.values():
        uteis_p = sorted(util_por[p.chave])
        p.dias_uteis = len(uteis_p)
        for d in uteis_p:
            linha = por_dia.get((p.chave, d), (0, {}))[1]
            c = calc.get((p.chave, d))
            do_csv = _batidas(linha)
            batidas = batidas_de.get((p.chave, d), do_csv)
            if (p.chave, d) in dias_fora:
                p.dias_fora += 1
                if batidas:
                    ev = dias_fora[(p.chave, d)]
                    pendencia(p, d, "aviso", f"marcação em dia de evento de dia inteiro ({ev.evidencia()}); "
                                             "dia fora do denominador — conferir")
                continue
            # 3. Evento que só o FLASH conhece: sem marcação, nada no CSV de
            # Eventos naquele dia (nem cancelado) e `Eventos` no Cálculo.
            if not batidas and c and c.tem_evento and (p.chave, d) not in com_evento:
                p.dias_fora += 1
                p.dias_fora_flash += 1
                pendencia(p, d, "aviso", f"FLASH: {c.eventos} (não consta no export de Eventos); "
                                         "dia fora do denominador")
                continue
            p.dias_elegiveis += 1
            evs = parciais.get((p.chave, d), [])
            nao_aprovada = [i for i in range(1, 7) if _minutos(linha.get(f"Hora da Marcação {i}", "")) is not None
                            and linha.get(f"Tipo Marcação {i}") == "Justificada"
                            and linha.get(f"Status Marcação {i}", "").strip().lower() != "aprovado"]
            if not batidas or nao_aprovada:
                p.dias_pendentes += 1
                for m in MEDIDAS:
                    p.pendentes[m] += 1
                dias.append(DiaAvaliado(p.chave, d, dict.fromkeys(MEDIDAS)))
                if nao_aprovada:
                    motivo = "batida justificada não aprovada — conferir no Flash"
                elif evs:
                    motivo = ("sem marcação; há evento parcial aprovado ("
                              + ", ".join(e.evidencia() for e in evs) + ") — conferir")
                elif c:
                    # O FLASH chama esses dias de `Falta`; nós continuamos não chamando.
                    motivo = (f"FLASH: {c.jornada}, eventos: {c.eventos} — sem marcação e sem "
                              "evento aprovado; não é falta, conferir")
                else:
                    motivo = "sem marcação e sem evento aprovado — não é falta; conferir"
                pendencia(p, d, "dia", motivo)
                continue

            n = len(batidas)
            medidas = {"entrada": (batidas[0], [1])}
            if n % 2 == 0:
                medidas["saida"] = (batidas[-1], [n])
            if n == 4:
                medidas["intervalo"] = (batidas[2] - batidas[1], [2, 3])
            dia_av = DiaAvaliado(p.chave, d, dict.fromkeys(MEDIDAS))
            texto_batidas = _texto_batidas(batidas)
            for m in MEDIDAS:
                if m not in medidas:
                    p.pendentes[m] += 1
                    medida_txt = "saída" if m == "saida" else "intervalo"
                    if c is None:
                        motivo = f"{n} batida(s): {medida_txt} não verificável — conferir no Cálculos Diários"
                    elif m == "intervalo" and n == 2:
                        # O FLASH não trata intervalo: conta o dia inteiro.
                        motivo = f"FLASH sem intervalo registrado (horas trabalhadas {c.horas})"
                    elif m == "intervalo" and n % 2 == 0:
                        # Qual das pausas é o intervalo é decisão do RH, não nossa.
                        motivo = _pausas(batidas)
                    else:
                        motivo = (f"{n} batida(s): {medida_txt} não verificável — FLASH: {c.jornada} "
                                  f"(horas trabalhadas {c.horas})")
                    pendencia(p, d, NOME_MEDIDA[m], motivo)
                    continue
                valor, posicoes = medidas[m]
                f = faixa(m, valor)
                dia_av.faixa[m] = f
                p.avaliados[m] += 1
                if f == NEUTRA:
                    p.neutras[m] += 1
                    continue
                quem = justificativa(f, batidas, evs)
                for e in quem:
                    e.justificou.append(f"{f} em {d:%d/%m}")
                situacao = JUSTIFICADA if quem else IDENTIFICADA
                (p.justificadas if quem else p.identificadas)[f] += 1
                if batidas != do_csv:
                    ajuste = f"valem as Marcações Válidas do FLASH (CSV de marcações: {_texto_batidas(do_csv)})"
                else:
                    ajuste = _evidencia_ajuste(linha, posicoes)
                ocorrencias.append(Ocorrencia(
                    p.chave, p.matricula, p.nome, d, f, valor, situacao,
                    "; ".join(e.evidencia() for e in quem), texto_batidas, ajuste,
                    c.jornada if c else "", c.eventos if c else ""))
            dias.append(dia_av)

    for e in eventos:
        if e.justificou:
            e.por_que += "; justificou " + ", ".join(e.justificou)

    if "banco" in lidos:
        for b in lidos["banco"]:
            dono = casar(_matricula(b.get("Matrícula", "")), b.get("Nome", ""))
            if dono:
                pessoas[dono].saldo_min, pessoas[dono].saldo_txt = _saldo(b.get("Saldo banco total", ""))

    ordem = {c: i for i, c in enumerate(sorted(pessoas, key=lambda c: (_sem_acento(pessoas[c].nome), c)))}
    r = Relatorio(
        mes=mes, ano=ano, arquivos=hashes, uteis=uteis,
        pessoas=sorted(pessoas.values(), key=lambda x: ordem[x.chave]),
        dias=dias,
        ocorrencias=sorted(ocorrencias, key=lambda o: (ordem[o.chave], o.data, o.indicador)),
        pendencias=sorted(pendencias, key=lambda x: (ordem[x.chave], x.data, x.medida)),
        eventos=eventos, tem_banco="banco" in lidos, tem_calculo=tem_calculo,
        avisos=sorted(avisos, key=lambda x: (ordem[x.chave], x.data, x.tipo)))
    if tem_calculo:
        r.calculo = {"trabalho": sum(1 for c in calc.values() if c.trabalho),
                     "marcacoes_comparadas": comparadas, "marcacoes_iguais": iguais,
                     "dias_fora_flash": sum(p.dias_fora_flash for p in pessoas.values())}
    r.conferencias = conferir(r)
    exigir_fechamento(r)
    return r


def _intervalo(inicio: date, fim: date):
    d = inicio
    while d <= fim:
        yield d
        d += timedelta(days=1)


def _texto_batidas(batidas: list[int]) -> str:
    return " · ".join(_fmt(b) for b in batidas)


def _batidas(linha: dict) -> list[int]:
    """`Hora da Marcação N` preenchidas, na ordem do Flash."""
    return [m for m in (_minutos(linha.get(f"Hora da Marcação {i}", "")) for i in range(1, 7))
            if m is not None]


# ── indicadores e fechamento ────────────────────────────────────────────────

def totais(r: Relatorio) -> dict[str, dict]:
    """O quadro A1–A6, contado das listas planas (dias e ocorrências) — NÃO
    da soma por pessoa. É o que dá sentido ao fechamento."""
    q = {}
    for a, (nome, medida) in INDICADORES.items():
        ident = sum(1 for o in r.ocorrencias if o.indicador == a and o.situacao == IDENTIFICADA)
        just = sum(1 for o in r.ocorrencias if o.indicador == a and o.situacao == JUSTIFICADA)
        denom = sum(1 for d in r.dias if d.faixa[medida] is not None)
        pend = sum(1 for d in r.dias if d.faixa[medida] is None)
        pessoas = len({o.chave for o in r.ocorrencias if o.indicador == a and o.situacao == IDENTIFICADA})
        q[a] = {"nome": nome, "identificadas": ident, "justificadas": just, "pendentes": pend,
                "denominador": denom, "pct": (ident / denom) if denom else None, "pessoas": pessoas}
    return q


def conferir(r: Relatorio) -> list[Conferencia]:
    """As somas do fechamento. Divergência é defeito da ferramenta."""
    q = totais(r)
    c: list[Conferencia] = []
    for a in INDICADORES:
        c.append(Conferencia(f"{a} identificadas: soma por colaborador × total",
                             q[a]["identificadas"], sum(p.identificadas[a] for p in r.pessoas)))
        c.append(Conferencia(f"{a} justificadas: soma por colaborador × total",
                             q[a]["justificadas"], sum(p.justificadas[a] for p in r.pessoas)))
    for m in MEDIDAS:
        baixo, rotulo, alto = FAIXAS[m]
        avaliados = sum(1 for d in r.dias if d.faixa[m] is not None)
        neutras = sum(1 for d in r.dias if d.faixa[m] == NEUTRA)
        soma = (q[baixo]["identificadas"] + q[baixo]["justificadas"] + neutras
                + q[alto]["identificadas"] + q[alto]["justificadas"])
        nm = NOME_MEDIDA[m]
        c.append(Conferencia(f"{nm}: {baixo} + {rotulo.lower()} + {alto} = avaliados", avaliados, soma))
        c.append(Conferencia(f"{nm}: avaliados por colaborador × total",
                             avaliados, sum(p.avaliados[m] for p in r.pessoas)))
        c.append(Conferencia(f"{nm}: pendentes por colaborador × total",
                             q[baixo]["pendentes"], sum(p.pendentes[m] for p in r.pessoas)))
        c.append(Conferencia(f"{nm}: avaliados + pendentes = dias elegíveis",
                             sum(p.dias_elegiveis for p in r.pessoas), avaliados + q[baixo]["pendentes"]))
    c.append(Conferencia("dias elegíveis = dias úteis − dias fora (soma por colaborador)",
                         sum(p.dias_uteis - p.dias_fora for p in r.pessoas),
                         sum(p.dias_elegiveis for p in r.pessoas)))
    c.append(Conferencia("dias pendentes: soma por colaborador × total",
                         sum(1 for d in r.dias if all(v is None for v in d.faixa.values())),
                         sum(p.dias_pendentes for p in r.pessoas)))
    return c


def exigir_fechamento(r: Relatorio) -> None:
    falhas = [x for x in conferir(r) if not x.ok]
    if falhas:
        detalhe = "; ".join(f"{x.verificacao}: esperado {x.esperado}, obtido {x.obtido}" for x in falhas)
        raise ErroPonto(f"A apuração não fecha ({detalhe}). É defeito da ferramenta, não dos "
                        "arquivos: não entrego planilha que não fecha. Avise o administrador.")


# ── histórico mensal (o JSON que o router grava no R2) ──────────────────────

def periodo(r: Relatorio) -> str:
    return f"{r.ano:04d}-{r.mes:02d}"


def periodo_anterior(per: str) -> str:
    ano, mes = (int(x) for x in per.split("-"))
    return f"{ano - 1:04d}-12" if mes == 1 else f"{ano:04d}-{mes - 1:02d}"


def historico(r: Relatorio, gerado_em: str) -> dict:
    """Só matrícula e números: sem nome, sem horário, sem evento.

    Quem não tem matrícula entra nos totais e fica fora de `por_matricula` —
    a chave dele seria o nome.
    """
    q = totais(r)
    por_mat = {}
    for p in r.pessoas:
        if not p.matricula:
            continue
        por_mat[p.matricula] = {
            "dias_uteis": p.dias_uteis, "dias_fora": p.dias_fora,
            "dias_elegiveis": p.dias_elegiveis, "dias_pendentes": p.dias_pendentes,
            **{a: {"identificadas": p.identificadas[a], "justificadas": p.justificadas[a],
                   "pendentes": p.pendentes[m], "denominador": p.avaliados[m]}
               for a, (_, m) in INDICADORES.items()},
        }
    return {
        "periodo": periodo(r), "versao_regua": VERSAO_REGUA, "gerado_em": gerado_em,
        "com_calculo_diario": r.tem_calculo,
        "arquivos": [{"nome": n, "sha256": h} for n, h in r.arquivos],
        "colaboradores": len(r.pessoas),
        "sem_matricula": sum(1 for p in r.pessoas if not p.matricula),
        "totais": {a: {k: v for k, v in t.items() if k not in ("nome", "pct")} for a, t in q.items()},
        "por_matricula": por_mat,
    }


def caminho_historico(h: dict) -> str:
    """`<AAAA-MM>/<versao>-<sha8 dos arquivos>.json` — dentro do bucket `ponto-historico`."""
    sha8 = hashlib.sha256("".join(sorted(a["sha256"] for a in h["arquivos"])).encode()).hexdigest()[:8]
    return f"{h['periodo']}/{h['versao_regua']}-{sha8}.json"


def _rotulo_periodo(per: str) -> str:
    ano, mes = (int(x) for x in per.split("-"))
    return f"{MESES[mes]}/{ano}"


def comparar(atual: dict, anteriores: list[dict]) -> Comparacao:
    """Compara com o mês anterior apurado na MESMA `VERSAO_REGUA` (o mais recente vale).

    Régua diferente não se compara: diz qual foi e pede reprocessamento. Com e
    sem o Cálculo Diário também não (v2.1): o calendário e os dias fora mudam
    com o arquivo, então é outra régua na prática, mesmo com a mesma versão.
    """
    alvo = periodo_anterior(atual["periodo"])
    rot = _rotulo_periodo(alvo)
    do_mes = [h for h in anteriores if isinstance(h, dict) and h.get("periodo") == alvo]
    mesma_versao = [h for h in do_mes if h.get("versao_regua") == atual["versao_regua"]]
    com_calc = bool(atual.get("com_calculo_diario"))
    mesma = [h for h in mesma_versao if bool(h.get("com_calculo_diario")) == com_calc]
    if mesma_versao and not mesma:
        antes, agora = ("sem", "com") if com_calc else ("com", "sem")
        return Comparacao("outra_versao", alvo,
                          f"{rot} foi apurado {antes} o Cálculo Diário; este mês, {agora}. "
                          f"Reprocesse {rot} {agora} ele para comparar.")
    if not mesma:
        if do_mes:
            versoes = ", ".join(sorted({str(h.get("versao_regua")) for h in do_mes}))
            return Comparacao("outra_versao", alvo,
                              f"{rot} foi apurado na régua {versoes}; esta é a {atual['versao_regua']}. "
                              f"Reprocesse {rot} com os arquivos dele para comparar.")
        return Comparacao("sem_historico", alvo, f"Não há apuração de {rot} no histórico; sem comparação.")
    ant = max(mesma, key=lambda h: str(h.get("gerado_em", "")))
    linhas = {}
    for a in INDICADORES:
        t_ant, t_at = ant["totais"][a], atual["totais"][a]
        pct_ant = t_ant["identificadas"] / t_ant["denominador"] if t_ant["denominador"] else None
        pct_at = t_at["identificadas"] / t_at["denominador"] if t_at["denominador"] else None
        linhas[a] = {"identificadas": t_ant["identificadas"], "denominador": t_ant["denominador"],
                     "pct": pct_ant,
                     "variacao_pp": None if pct_ant is None or pct_at is None else (pct_at - pct_ant) * 100}
    return Comparacao("comparado", alvo, f"Comparado com {rot} (mesma régua, {atual['versao_regua']}).",
                      linhas)


# ── saída: xlsx e resumo ────────────────────────────────────────────────────

MESES = ("", "JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO")
ABAS = ("Resumo", "Por colaborador", "Ocorrências", "Pendências", "Eventos", "Conferência")


def nome_do_arquivo(r: Relatorio) -> str:
    return f"Controle de jornada - {MESES[r.mes]}-{r.ano}.xlsx"


def rotulo(a: str) -> str:
    return f"{a} — {INDICADORES[a][0]}"


def _hora(minutos: int | None) -> float | None:
    return None if minutos is None else minutos / 1440


def _pct(x: float | None) -> str:
    return "—" if x is None else f"{x * 100:.1f}%".replace(".", ",")


def _cabecalho(ws, colunas: list[str]) -> None:
    from openpyxl.styles import Font
    ws.append(colunas)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)


def _formatar(ws, colunas: dict[int, str]) -> None:
    for col, fmt in colunas.items():
        ws.cell(ws.max_row, col).number_format = fmt


def montar_xlsx(r: Relatorio) -> bytes:
    """Números como número; horário como fração de dia em `hh:mm`; % como fração."""
    import openpyxl

    wb = openpyxl.Workbook()
    q = totais(r)

    # 1. Resumo
    ws = wb.active
    ws.title = ABAS[0]
    ws.append([f"CONTROLE DE JORNADA — {MESES[r.mes]}/{r.ano}"])
    ws.append(["Grupo A — Pontualidade e Jornada. Critérios gerenciais: não alteram saldo do Flash."])
    ws.append([])
    _cabecalho(ws, ["Período", "Versão da régua", "Colaboradores", "Dias úteis do mês"])
    ws.append([f"{r.mes:02d}/{r.ano}", VERSAO_REGUA, len(r.pessoas), len(r.uteis)])
    ws.append([])
    _cabecalho(ws, ["Arquivo processado", "sha256"])
    for nome, h in r.arquivos:
        ws.append([nome, h[:12]])
    ws.append([])
    comp = r.comparacao
    extra = (["Mês anterior: identificadas", "Mês anterior: denominador", "Mês anterior: %",
              "Variação (p.p.)"] if comp and comp.status == "comparado" else [])
    _cabecalho(ws, ["Indicador", "Identificadas", "Justificadas", "Pendentes", "Denominador", "%",
                    "Colaboradores envolvidos", *extra])
    for a, t in q.items():
        linha = [rotulo(a), t["identificadas"], t["justificadas"], t["pendentes"], t["denominador"],
                 t["pct"], t["pessoas"]]
        if extra:
            ant = comp.linhas[a]
            linha += [ant["identificadas"], ant["denominador"], ant["pct"],
                      None if ant["variacao_pp"] is None else round(ant["variacao_pp"], 2)]
        ws.append(linha)
        _formatar(ws, {6: "0.0%", **({10: "0.0%"} if extra else {})})
    ws.append([])
    _cabecalho(ws, ["Faixa sem ocorrência", "Dias"])
    for m in MEDIDAS:
        ws.append([FAIXAS[m][1], sum(1 for d in r.dias if d.faixa[m] == NEUTRA)])
    ws.append([])
    usados = sum(1 for e in r.eventos if e.usado)
    _cabecalho(ws, ["Outros números", "Valor"])
    for rot, valor in (
        ("Dias elegíveis (soma)", sum(p.dias_elegiveis for p in r.pessoas)),
        ("Dias fora por evento de dia inteiro (soma)", sum(p.dias_fora for p in r.pessoas)),
        ("Dias pendentes (soma)", sum(p.dias_pendentes for p in r.pessoas)),
        ("Pendências listadas", len(r.pendencias)),
        ("Eventos usados", usados),
        ("Eventos não usados", len(r.eventos) - usados),
        ("Fechamento", "conferido" if all(c.ok for c in r.conferencias) else "NÃO FECHA"),
    ):
        ws.append([rot, valor])
    ws.append([])
    if r.tem_calculo:
        k = r.calculo
        ws.append(["Cálculo Diário do FLASH", "usado: calendário e marcações do FLASH"])
        ws.append(["Marcações iguais ao CSV (dias)", k["marcacoes_iguais"]])
        ws.append(["Marcações comparadas (dias)", k["marcacoes_comparadas"]])
        ws.append(["Pessoa-dia de Trabalho no FLASH", k["trabalho"]])
        ws.append(["Dias fora só pelo FLASH (sem evento no export)", k["dias_fora_flash"]])
        ws.append(["Avisos (aba Conferência)", len(r.avisos)])
    else:
        ws.append(["Cálculo Diário do FLASH", "não veio: calendário pela lista de feriados; "
                                              "dias fora do padrão sem a evidência do FLASH"])
    if comp:
        ws.append([])
        ws.append(["Comparação com o mês anterior", comp.mensagem])
    ws.append([])
    ws.append(["Rótulos de reincidência (crítico, atenção, acompanhamento) estão suspensos "
               "pelo RH até validação de critério novo."])

    # 2. Por colaborador — sem rótulo de status.
    ws = wb.create_sheet(ABAS[1])
    cols = ["Matrícula", "Nome", "Setor", "Dias úteis", "Dias fora (evento)", "Dias elegíveis",
            "Dias pendentes"]
    for a in INDICADORES:
        cols += [f"{a} identificadas", f"{a} denominador", f"{a} %", f"{a} justificadas"]
    cols += ["Saldo banco", "Saldo (min)"]
    if r.tem_calculo:
        cols += ["Dias fora só pelo FLASH"]
    _cabecalho(ws, cols)
    for p in r.pessoas:
        linha = [p.matricula, p.nome, p.setor, p.dias_uteis, p.dias_fora, p.dias_elegiveis,
                 p.dias_pendentes]
        fmts = {}
        for a, (_, m) in INDICADORES.items():
            den = p.avaliados[m]
            linha += [p.identificadas[a], den, (p.identificadas[a] / den) if den else None,
                      p.justificadas[a]]
            fmts[len(linha) - 1] = "0.0%"
        linha += [p.saldo_txt, p.saldo_min]
        if r.tem_calculo:
            linha += [p.dias_fora_flash]
        ws.append(linha)
        _formatar(ws, fmts)

    # 3. Ocorrências — a evidência do FLASH vai no fim, só quando o Cálculo veio.
    ws = wb.create_sheet(ABAS[2])
    _cabecalho(ws, ["Matrícula", "Nome", "Data", "Indicador", "Horário / duração", "Situação",
                    "Evento que justifica", "Batidas do dia", "Batida ajustada"]
               + (["FLASH: jornada considerada", "FLASH: eventos"] if r.tem_calculo else []))
    for o in r.ocorrencias:
        ws.append([o.matricula, o.nome, o.data, rotulo(o.indicador), _hora(o.minutos), o.situacao,
                   o.evento, o.batidas, o.ajuste]
                  + ([o.flash_jornada, o.flash_eventos] if r.tem_calculo else []))
        _formatar(ws, {3: "dd/mm/yyyy", 5: "hh:mm"})

    # 4. Pendências
    ws = wb.create_sheet(ABAS[3])
    _cabecalho(ws, ["Matrícula", "Nome", "Data", "Medida", "Motivo"]
               + (["FLASH: jornada considerada", "FLASH: horas trabalhadas", "FLASH: marcações válidas",
                   "FLASH: eventos"] if r.tem_calculo else []))
    for x in r.pendencias:
        ws.append([x.matricula, x.nome, x.data, x.medida, x.motivo]
                  + ([x.flash_jornada, x.flash_horas, x.flash_marcacoes, x.flash_eventos]
                     if r.tem_calculo else []))
        _formatar(ws, {3: "dd/mm/yyyy"})

    # 5. Eventos — sem CID e sem justificativa (nem chegaram a ser lidos).
    ws = wb.create_sheet(ABAS[4])
    _cabecalho(ws, ["EventoId", "Matrícula", "Nome", "Motivo", "Código", "Início", "Fim",
                    "Hora início", "Hora fim", "Qtde", "Status", "Usado", "Por quê"])
    for e in sorted(r.eventos, key=lambda e: (e.inicio, _sem_acento(e.nome), e.id)):
        ws.append([e.id, e.matricula, e.nome, e.motivo, e.codigo, e.inicio, e.fim,
                   _hora(e.hora_inicio), _hora(e.hora_fim), e.qtde, e.status,
                   "sim" if e.usado else "não", e.por_que])
        _formatar(ws, {6: "dd/mm/yyyy", 7: "dd/mm/yyyy", 8: "hh:mm", 9: "hh:mm"})

    # 6. Conferência
    ws = wb.create_sheet(ABAS[5])
    _cabecalho(ws, ["Verificação", "Esperado", "Obtido", "Fecha?"])
    for c in r.conferencias:
        ws.append([c.verificacao, c.esperado, c.obtido, "SIM" if c.ok else "NÃO"])
    if r.tem_calculo:
        # Avisos não são fechamento: a apuração sai, e a RH confere a diferença.
        ws.append([])
        _cabecalho(ws, ["Cálculo Diário × demais arquivos", "Valor"])
        ws.append(["Marcações Válidas iguais às batidas do CSV (dias)",
                   f"{r.calculo['marcacoes_iguais']} de {r.calculo['marcacoes_comparadas']}"])
        ws.append(["Pessoa-dia de Trabalho no FLASH", r.calculo["trabalho"]])
        ws.append([])
        _cabecalho(ws, ["Aviso", "Matrícula", "Nome", "Data", "Detalhe"])
        for a in r.avisos:
            ws.append([a.tipo, a.matricula, a.nome, a.data, a.texto])
            _formatar(ws, {4: "dd/mm/yyyy"})
        if not r.avisos:
            ws.append(["nenhum: calendário, marcações e pessoa-dia batem"])

    # Injeção de fórmula: nome, motivo e setor são texto que vem do Flash.
    # Começando com = + - @, o Excel da RH os executaria (ex.: =HYPERLINK).
    # Força tipo texto — o mesmo do `enviar_arquivo` em `anexos_escrita.py`.
    for planilha in wb.worksheets:
        for linha in planilha.iter_rows():
            for c in linha:
                if isinstance(c.value, str) and c.value[:1] in ("=", "+", "-", "@"):
                    c.data_type = "s"

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def resumo(r: Relatorio) -> str:
    """O que volta ao modelo: o quadro e os totais — NUNCA nome de pessoa.

    Os detalhes por colaborador são da planilha, que é privada; o resumo vai
    para o contexto do agente, que pode citá-lo em qualquer lugar.
    """
    q = totais(r)
    usados = sum(1 for e in r.eventos if e.usado)
    nao_usados = Counter(e.por_que.split(" — ")[0].removeprefix("não: ")
                         for e in r.eventos if not e.usado)
    partes = [
        f"Controle de jornada {MESES[r.mes]}/{r.ano} — régua {VERSAO_REGUA} "
        "(critérios gerenciais; não mexem em saldo do Flash).",
        f"{len(r.pessoas)} colaboradores, {len(r.uteis)} dias úteis no mês; "
        f"{sum(p.dias_elegiveis for p in r.pessoas)} dias elegíveis, "
        f"{sum(p.dias_fora for p in r.pessoas)} fora por evento de dia inteiro, "
        f"{sum(p.dias_pendentes for p in r.pessoas)} dias pendentes.",
        "Quadro A1–A6 (identificadas / denominador = %; justificadas; pendentes; colaboradores):",
    ]
    for a, t in q.items():
        partes.append(f"- {rotulo(a)}: {t['identificadas']}/{t['denominador']} = {_pct(t['pct'])}; "
                      f"{t['justificadas']} justificada(s); {t['pendentes']} pendente(s); "
                      f"{t['pessoas']} colaborador(es)")
    por_motivo = Counter(x.medida for x in r.pendencias)
    partes.append(f"Pendências: {len(r.pendencias)} ("
                  + ", ".join(f"{k}: {v}" for k, v in sorted(por_motivo.items())) + ")."
                  if r.pendencias else "Pendências: nenhuma.")
    por_status = Counter((e.status or "vazio").lower() for e in r.eventos)
    partes.append(f"Eventos no arquivo: {len(r.eventos)} ("
                  + ", ".join(f"{v} {k}" for k, v in sorted(por_status.items())) + f"); "
                  f"{usados} usados, {len(r.eventos) - usados} não usados"
                  + (" — " + "; ".join(f"{v} {k}" for k, v in sorted(nao_usados.items()))
                     if nao_usados else "") + ".")
    partes.append("Fechamento: conferido — as somas por colaborador batem com o total."
                  if all(c.ok for c in r.conferencias) else "Fechamento: NÃO FECHA.")
    if r.comparacao:
        partes.append(r.comparacao.mensagem)
        if r.comparacao.status == "comparado":
            for a, ant in r.comparacao.linhas.items():
                var = ant["variacao_pp"]
                partes.append(f"  {a}: antes {ant['identificadas']}/{ant['denominador']} = {_pct(ant['pct'])}"
                              + ("" if var is None else f", variação {var:+.1f} p.p.".replace(".", ",")))
    if r.tem_calculo:
        k = r.calculo
        por_tipo = Counter(a.tipo for a in r.avisos)
        partes.append(
            f"Cálculo Diário: marcações iguais em {k['marcacoes_iguais']} de {k['marcacoes_comparadas']} "
            f"dias; {k['trabalho']} pessoa-dia de Trabalho no FLASH; {k['dias_fora_flash']} dia(s) fora "
            "só pelo FLASH (evento que não consta no export de Eventos); avisos: "
            + (", ".join(f"{k2}: {v}" for k2, v in sorted(por_tipo.items())) if por_tipo else "nenhum")
            + ".")
    else:
        partes.append("Sem o CSV de Cálculo Diário (CalculoDiario): calendário pela lista de "
                      "feriados da régua e dias fora do padrão sem a evidência do FLASH. É "
                      "opcional, mas recomendado — peça na próxima vez.")
    if not r.tem_banco:
        partes.append("Sem o CSV de banco de horas: a planilha sai sem saldo.")
    return "\n".join(partes)
