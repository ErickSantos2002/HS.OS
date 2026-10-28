# backend/tests/test_relatorio_ponto.py
"""A régua de ponto v2 (Grupo A) — a da coordenadora de RH, não a nossa.

Mandada por escrito em 27/10/2026, depois de duas semanas homologando o
`bruce` contra os relatórios do Flash. Ver
`docs/superpowers/specs/2026-10-28-regua-ponto-v2-design.md`.

⚠️ **A lição que justifica o módulo inteiro:** na homologação o `bruce` não
usou a ferramenta (ela aceitava dois arquivos e vieram três), leu os CSV com
`ler_anexo` e contou de cabeça — 21 eventos aprovados e 2 cancelados, quando o
arquivo tinha 20 e 3. E inventou um "deslocamento de coluna" para explicar a
própria leitura. Número sai de código, com teste na fronteira.

⚠️ **Os dados aqui são inventados.** O repositório é público e os CSV reais
trazem nome, matrícula, horário e diagnóstico de cada funcionário.
"""
import io as _io
import json
from datetime import date, time

import openpyxl
import pytest

from app.relatorios import ponto as p

CAB_MARC = ["Nome", "Matrícula", "Filial", "Departamento", "Centro de Custo", "Data da Marcação"]
for _i in range(1, 7):
    CAB_MARC += [f"Tipo Marcação {_i}", f"Motivo Marcação {_i}", f"Justificativa {_i}",
                 f"Status Marcação {_i}", f"Hora da Marcação {_i}"]
CAB_BANCO = ["Ano", "Mês", "Matrícula", "Nome", "Cargo", "Filial", "Nome Departamento",
             "Código Centro de Custo", "Nome Gestor", "Saldo banco anterior",
             "Saldo banco no periodo", "Saldo banco total"]
CAB_EVENTOS = ["Matrícula", "Nome", "Documento", "Departamento", "Filial", "EventoId",
               "Data da Inclusão", "Data da Aprovação", "Aprovador", "Observação da Aprovação",
               "Motivo", "Descrição", "Justificativa", "CID", "Data Inicio", "Data Fim",
               "Hora Inicio", "Hora Fim", "Qtde", "Status", "Anexos", "FilialCNPJ", "Cargo",
               "Status do Funcionário"]

ANA, BIA = ("Ana Teste", "000001"), ("Bia Teste", "000002")
NORMAL = ("07:30", "12:00", "13:00", "17:00")
# Setembro/2026: 22 dias de semana, menos 07/09 (Independência) = 21 dias úteis.
UTEIS_SET = [1, 2, 3, 4, 8, 9, 10, 11, 14, 15, 16, 17, 18, 21, 22, 23, 24, 25, 28, 29, 30]


def marc(pessoa, dia, *batidas, depto="LABORATÓRIO"):
    """Uma linha de marcações. `dia` é o dia de setembro/2026 (int) ou "dd/mm/aaaa".
    Cada batida é "HH:MM" ou (hora, tipo, motivo, justificativa[, status])."""
    nome, matricula = pessoa
    data = f"{dia:02d}/09/2026" if isinstance(dia, int) else dia
    linha = [nome, matricula, "MATRIZ", depto, "", data]
    for b in batidas:
        if isinstance(b, str):
            b = (b, "Original", "", "")
        hora, tipo, motivo, just = b[:4]
        status = b[4] if len(b) > 4 else ("Aprovado" if tipo == "Justificada" else "")
        linha += [tipo, motivo, just, status, hora]
    linha += [""] * (len(CAB_MARC) - len(linha))
    return linha


def mes_inteiro(pessoa, batidas=NORMAL, exceto=()):
    return [marc(pessoa, d, *batidas) for d in UTEIS_SET if d not in exceto]


def ev(pessoa, inicio, fim="", hi="", hf="", qtde="", status="Aprovado", motivo="Declaração de Horas",
       ident="9000001", descricao="", cid="", justificativa=""):
    nome, matricula = pessoa
    qtde = qtde or ("Dia Todo" if not hi else "01:00")
    return [matricula, nome, "000.000.000-00", "LABORATÓRIO", "MATRIZ", ident, "01/09/2026",
            "02/09/2026", "GESTOR FICTICIO", "", motivo, descricao, justificativa, cid,
            inicio, fim, hi, hf, qtde, status, "", "", "Analista", "Ativo"]


def csv(cab, linhas, bom=True, codif="utf-8"):
    texto = "\n".join(";".join(l) for l in [cab] + linhas) + "\n"
    return (("﻿" if bom else "") + texto).encode(codif)


def banco(*pessoas):
    """pessoas: (nome, matricula, saldo_total)."""
    return csv(CAB_BANCO, [["2026", "9", m, n, "Analista", "MATRIZ", "LABORATÓRIO", "",
                            "GESTOR FICTICIO", "00:00", "00:00", s] for n, m, s in pessoas])


def rel(linhas_marc, eventos=(), pessoas_banco=None):
    arquivos = [("m.csv", csv(CAB_MARC, linhas_marc)), ("e.csv", csv(CAB_EVENTOS, list(eventos)))]
    if pessoas_banco is not None:
        arquivos.append(("b.csv", banco(*pessoas_banco)))
    return p.gerar(arquivos)


def pessoa(r, nome="Ana Teste"):
    return next(x for x in r.pessoas if x.nome == nome)


def ocorr(r, dia=None):
    return [(o.indicador, o.situacao) for o in r.ocorrencias if dia is None or o.data.day == dia]


def _livro(r):
    return openpyxl.load_workbook(_io.BytesIO(p.montar_xlsx(r)))


def _todo_texto_do_xlsx(r):
    wb = _livro(r)
    return "\n".join(str(c.value) for ws in wb.worksheets for linha in ws.iter_rows()
                     for c in linha if c.value is not None)


# ── leitura ─────────────────────────────────────────────────────────────────

def test_reconhece_os_tres_arquivos_pelo_cabecalho_em_qualquer_ordem():
    m = csv(CAB_MARC, [marc(ANA, 1, *NORMAL)])
    e = csv(CAB_EVENTOS, [])
    b = banco((*ANA, "01:00"))
    assert p.identificar("x.csv", m)[0] == "marcacoes"
    assert p.identificar("y.csv", e)[0] == "eventos"
    assert p.identificar("z.csv", b)[0] == "banco"
    r = p.gerar([("z.csv", b), ("y.csv", e), ("x.csv", m)])
    assert (r.mes, r.ano) == (9, 2026) and len(r.pessoas) == 1 and r.tem_banco


def test_banco_e_opcional():
    r = rel([marc(ANA, 1, *NORMAL)])
    assert not r.tem_banco and pessoa(r).saldo_min is None
    assert "sem saldo" in p.resumo(r)


def test_sem_eventos_e_recusado_e_a_mensagem_nomeia_o_arquivo_que_falta():
    """Sem eventos não há denominador: férias e atestado virariam "dia sem
    marcação" e o percentual sairia errado sem ninguém perceber."""
    m = csv(CAB_MARC, [marc(ANA, 1, *NORMAL)])
    with pytest.raises(p.ErroPonto, match=r"eventos \(Eventos\).*obrigatório"):
        p.gerar([("m.csv", m), ("b.csv", banco((*ANA, "01:00")))])


def test_sem_marcacoes_e_recusado():
    with pytest.raises(p.ErroPonto, match="marcações"):
        p.gerar([("e.csv", csv(CAB_EVENTOS, []))])


def test_dois_arquivos_do_mesmo_tipo_sao_recusados():
    m = csv(CAB_MARC, [marc(ANA, 1, *NORMAL)])
    with pytest.raises(p.ErroPonto, match="dois CSV de marcações"):
        p.gerar([("a.csv", m), ("b.csv", m), ("e.csv", csv(CAB_EVENTOS, []))])


def test_cabecalho_desconhecido_nomeia_o_arquivo():
    with pytest.raises(p.ErroPonto, match="planilha.csv"):
        p.identificar("planilha.csv", b"a;b;c\n1;2;3\n")


def test_le_csv_sem_bom_e_em_latin1():
    linhas = [marc(("Joana Ação", "000001"), 1, *NORMAL)]
    for dados in (csv(CAB_MARC, linhas, bom=False), csv(CAB_MARC, linhas, bom=False, codif="latin-1")):
        tipo, registros = p.identificar("m.csv", dados)
        assert tipo == "marcacoes" and registros[0]["Nome"] == "Joana Ação"


def test_mais_de_um_mes_e_recusado():
    with pytest.raises(p.ErroPonto, match="09/2026.*10/2026|10/2026.*09/2026"):
        rel([marc(ANA, 30, "07:30"), marc(ANA, "01/10/2026", "07:30")])


def test_teto_de_linhas():
    linhas = [marc(ANA, 1, "07:30")] * (p.TETO_LINHAS + 1)
    with pytest.raises(p.ErroPonto, match="linhas"):
        rel(linhas)


def test_data_invalida_nomeia_o_arquivo_e_a_linha():
    for data in ("2026-09-01", "", "01/09/26"):
        with pytest.raises(p.ErroPonto, match=r"m\.csv.*linha 2"):
            rel([marc(ANA, data, "07:30")])


def test_mesma_pessoa_e_dia_em_duas_linhas_e_recusado():
    """Contar o dia duas vezes infla numerador e denominador em silêncio."""
    with pytest.raises(p.ErroPonto, match="linha 3 repete"):
        rel([marc(ANA, 1, *NORMAL), marc(ANA, 1, *NORMAL)])


# ── calendário ──────────────────────────────────────────────────────────────

def test_feriados_de_2026_que_ela_deu():
    f = p.feriados(2026)
    assert f[date(2026, 9, 7)][0] == "Independência"
    assert f[date(2026, 2, 17)][0] == "Carnaval"
    assert f[date(2026, 4, 3)][0] == "Paixão de Cristo"
    assert len(f) == 15


@pytest.mark.parametrize("ano,pascoa,carnaval,paixao", [
    (2024, date(2024, 3, 31), date(2024, 2, 13), date(2024, 3, 29)),
    (2025, date(2025, 4, 20), date(2025, 3, 4), date(2025, 4, 18)),
    (2027, date(2027, 3, 28), date(2027, 2, 9), date(2027, 3, 26)),
])
def test_carnaval_e_paixao_saem_da_pascoa_em_outros_anos(ano, pascoa, carnaval, paixao):
    """Fixar 17/02 e 03/04 quebraria a régua em janeiro de 2027."""
    assert p.pascoa(ano) == pascoa
    f = p.feriados(ano)
    assert f[carnaval][0] == "Carnaval" and f[paixao][0] == "Paixão de Cristo"


def test_setembro_de_2026_tem_21_dias_uteis():
    uteis = p.dias_uteis(2026, 9)
    assert [d.day for d in uteis] == UTEIS_SET
    assert date(2026, 9, 7) not in uteis


def test_marcacao_no_sabado_vira_pendencia_fora_da_escala():
    r = rel([marc(ANA, 5, "06:00", "10:00")])   # 05/09/2026 é sábado
    assert ocorr(r) == []
    [x] = [x for x in r.pendencias if x.medida == "fora da escala"]
    assert x.data == date(2026, 9, 5) and "sábado" in x.motivo
    assert all(d.data != date(2026, 9, 5) for d in r.dias)


def test_marcacao_no_feriado_vira_pendencia_com_o_nome_do_feriado():
    r = rel([marc(ANA, 7, "07:00", "12:00", "13:00", "18:00")])
    [x] = [x for x in r.pendencias if x.medida == "fora da escala"]
    assert "Independência" in x.motivo and ocorr(r) == []


# ── a régua, na fronteira exata ─────────────────────────────────────────────

@pytest.mark.parametrize("entrada,esperado", [
    ("07:19", ["A1"]), ("07:19:59", ["A1"]), ("07:20", []), ("07:45", []),
    ("07:45:59", []), ("07:46", ["A2"]),
])
def test_limites_da_entrada(entrada, esperado):
    r = rel([marc(ANA, 1, entrada, "12:00", "13:00", "17:00")])
    assert [i for i, _ in ocorr(r)] == esperado


@pytest.mark.parametrize("saida,esperado", [
    ("16:49", ["A3"]), ("16:50", []), ("17:10", []), ("17:10:59", []), ("17:11", ["A4"]),
])
def test_limites_da_saida(saida, esperado):
    r = rel([marc(ANA, 1, "07:30", "12:00", "13:00", saida)])
    assert [i for i, _ in ocorr(r)] == esperado


@pytest.mark.parametrize("volta,esperado", [("12:59", ["A5"]), ("13:00", []), ("13:01", ["A6"])])
def test_intervalo_59_60_61(volta, esperado):
    r = rel([marc(ANA, 1, "07:30", "12:00", volta, "17:00")])
    assert [i for i, _ in ocorr(r)] == esperado


def test_ocorrencia_traz_horario_situacao_e_todas_as_batidas():
    r = rel([marc(ANA, 2, "07:52", "12:00", "13:00", "17:05")])
    [o] = r.ocorrencias
    assert (o.indicador, o.minutos, o.situacao) == ("A2", 7 * 60 + 52, "identificada")
    assert o.batidas == "07:52 · 12:00 · 13:00 · 17:05" and o.data == date(2026, 9, 2)


# ── batidas fora do padrão ──────────────────────────────────────────────────

@pytest.mark.parametrize("batidas", [("07:50", "17:20"),
                                     ("07:50", "12:00", "13:00", "15:00", "15:30", "17:20")])
def test_2_e_6_batidas_avaliam_entrada_e_saida_e_deixam_o_intervalo_pendente(batidas):
    """Ela foi explícita: não excluir o dia e não presumir intervalo. A v1
    pegava 3ª − 2ª nos dias de 6 batidas, o que mede a saída do meio da tarde."""
    r = rel([marc(ANA, 1, *batidas)])
    assert sorted(i for i, _ in ocorr(r)) == ["A2", "A4"]
    a = pessoa(r)
    assert a.avaliados == {"entrada": 1, "saida": 1, "intervalo": 0}
    [x] = [x for x in r.pendencias if x.data.day == 1]
    assert x.medida == "intervalo" and "Cálculos Diários" in x.motivo


def test_3_batidas_deixam_saida_e_intervalo_pendentes():
    r = rel([marc(ANA, 1, "07:50", "12:00", "13:00")])
    assert ocorr(r) == [("A2", "identificada")]
    assert sorted(x.medida for x in r.pendencias if x.data.day == 1) == ["intervalo", "saída"]


def test_dia_sem_marcacao_e_sem_evento_e_pendente_nao_ocorrencia():
    """Não é falta: pode ser admissão, desligamento ou export incompleto. Sai
    do numerador e do denominador e vai para a aba de pendências."""
    r = rel(mes_inteiro(ANA, exceto=(15,)))
    a = pessoa(r)
    assert ocorr(r) == []
    assert (a.dias_uteis, a.dias_fora, a.dias_elegiveis, a.dias_pendentes) == (21, 0, 21, 1)
    assert a.avaliados == {"entrada": 20, "saida": 20, "intervalo": 20}
    [x] = r.pendencias
    assert x.data.day == 15 and x.medida == "dia" and "não é falta" in x.motivo


def test_batida_justificada_aprovada_vale_e_vai_para_a_evidencia():
    r = rel([marc(ANA, 1, ("07:55", "Justificada", "Esquecimento de marcação", "esqueci"),
                  "12:00", "13:00", "17:00")])
    [o] = r.ocorrencias
    assert o.indicador == "A2" and o.ajuste == "batida ajustada: Esquecimento de marcação"


def test_batida_justificada_nao_aprovada_deixa_o_dia_inteiro_pendente():
    r = rel([marc(ANA, 1, ("07:55", "Justificada", "Esquecimento", "", "Pendente"),
                  "12:00", "13:00", "17:00")])
    assert ocorr(r) == []
    [x] = [x for x in r.pendencias if x.data.day == 1]
    assert x.medida == "dia" and "não aprovada" in x.motivo


# ── eventos ─────────────────────────────────────────────────────────────────

def test_evento_cancelado_nao_justifica_e_aparece_como_nao_usado():
    r = rel([marc(ANA, 1, "09:00", "12:00", "13:00", "17:00")],
            [ev(ANA, "01/09/2026", "01/09/2026", "07:30", "09:00", status="Cancelado")])
    assert ocorr(r) == [("A2", "identificada")]
    [e] = r.eventos
    assert not e.usado and "Cancelado" in e.por_que


def test_cancelado_de_dia_inteiro_nao_tira_o_dia():
    r = rel(mes_inteiro(ANA), [ev(ANA, "02/09/2026", status="Cancelado")])
    assert pessoa(r).dias_fora == 0 and pessoa(r).dias_elegiveis == 21


def test_evento_de_dia_inteiro_tira_o_dia_do_denominador():
    r = rel(mes_inteiro(ANA, exceto=(2,)), [ev(ANA, "02/09/2026", motivo="Atestado Médico")])
    a = pessoa(r)
    assert (a.dias_fora, a.dias_elegiveis, a.dias_pendentes) == (1, 20, 0)
    assert a.avaliados["entrada"] == 20 and r.pendencias == []
    assert r.eventos[0].usado


def test_marcacao_em_dia_de_evento_de_dia_inteiro_nao_conta_e_vira_aviso():
    r = rel([marc(ANA, 2, "09:00", "12:00", "13:00", "17:00")], [ev(ANA, "02/09/2026", motivo="Day-off")])
    assert ocorr(r) == [] and pessoa(r).dias_fora == 1
    [x] = [x for x in r.pendencias if x.data.day == 2]
    assert x.medida == "aviso" and "Day-off" in x.motivo


def test_evento_de_varios_dias_conta_so_os_uteis():
    """04/09 (sex) a 08/09 (ter): sábado, domingo e o feriado de 07/09 não são
    dia útil — saem 2 dias, não 5."""
    r = rel(mes_inteiro(ANA, exceto=(4, 8)),
            [ev(ANA, "04/09/2026", "08/09/2026", qtde="5 Dias", motivo="Férias")])
    assert pessoa(r).dias_fora == 2 and pessoa(r).dias_pendentes == 0


@pytest.mark.parametrize("ini,fim,situacao", [
    ("07:30", "10:00", "justificada"),   # começa até 07:45 e cobre a entrada
    ("07:45", "09:50", "justificada"),   # limites inclusivos
    ("07:46", "10:00", "identificada"),  # começa depois de 07:45
    ("07:30", "09:00", "identificada"),  # acaba antes da entrada
])
def test_evento_parcial_justifica_atraso(ini, fim, situacao):
    r = rel([marc(ANA, 1, "09:50", "12:00", "13:00", "17:00")],
            [ev(ANA, "01/09/2026", "01/09/2026", ini, fim)])
    assert ocorr(r) == [("A2", situacao)]


@pytest.mark.parametrize("ini,fim,situacao", [
    ("15:00", "17:00", "justificada"),
    ("15:05", "16:50", "justificada"),
    ("15:06", "17:00", "identificada"),  # começa depois da saída
    ("15:00", "16:49", "identificada"),  # acaba antes de 16:50
])
def test_evento_parcial_justifica_saida_antecipada(ini, fim, situacao):
    r = rel([marc(ANA, 1, "07:30", "12:00", "13:00", "15:05")],
            [ev(ANA, "01/09/2026", "01/09/2026", ini, fim)])
    assert ocorr(r) == [("A3", situacao)]


@pytest.mark.parametrize("ini,fim,situacao", [
    ("13:00", "14:00", "justificada"),   # 120 − 60 = 60 ≤ 60
    ("11:00", "13:00", "justificada"),   # sobreposição 60 (12:00–13:00)
    ("13:30", "14:00", "identificada"),  # 120 − 30 = 90 > 60
    ("15:00", "16:00", "identificada"),  # não sobrepõe o intervalo
])
def test_evento_parcial_justifica_intervalo_prolongado(ini, fim, situacao):
    r = rel([marc(ANA, 1, "07:30", "12:00", "14:00", "17:00")],
            [ev(ANA, "01/09/2026", "01/09/2026", ini, fim)])
    assert ocorr(r) == [("A6", situacao)]


def test_evento_parcial_nao_justifica_a1_a4_a5():
    """Chegar cedo, sair tarde e voltar antes do intervalo não são ausência."""
    r = rel([marc(ANA, 1, "07:00", "12:00", "12:30", "17:30")],
            [ev(ANA, "01/09/2026", "01/09/2026", "06:00", "18:00")])
    assert sorted(ocorr(r)) == [("A1", "identificada"), ("A4", "identificada"), ("A5", "identificada")]


def test_justificada_fica_fora_do_percentual_mas_aparece_ao_lado():
    r = rel([marc(ANA, 1, "09:50", "12:00", "13:00", "17:00"),
             marc(ANA, 2, "07:50", "12:00", "13:00", "17:00")],
            [ev(ANA, "01/09/2026", "01/09/2026", "07:30", "10:00",
                descricao="As horas vão abonar as horas negativas do dia (1004)")])
    t = p.totais(r)["A2"]
    assert (t["identificadas"], t["justificadas"], t["denominador"], t["pessoas"]) == (1, 1, 2, 1)
    assert t["pct"] == 0.5
    [j] = [o for o in r.ocorrencias if o.situacao == "justificada"]
    assert "(1004)" in j.evento and "9000001" in j.evento
    assert "justificou A2 em 01/09" in r.eventos[0].por_que


def test_evento_casa_por_nome_quando_falta_matricula():
    r = rel(mes_inteiro(("José Sem Matrícula", "-"), exceto=(3,)),
            [ev(("JOSE SEM MATRICULA", "-"), "03/09/2026")])
    a = pessoa(r, "José Sem Matrícula")
    assert a.dias_fora == 1 and a.dias_pendentes == 0


def test_evento_de_quem_nao_tem_marcacao_nao_e_usado():
    """No setembro real: pessoa inativa, com evento aprovado e sem marcação."""
    r = rel([marc(ANA, 1, *NORMAL)], [ev(BIA, "02/09/2026")])
    [e] = r.eventos
    assert not e.usado and "sem marcação" in e.por_que
    assert [x.nome for x in r.pessoas] == ["Ana Teste"]


def test_contagem_de_eventos_sai_do_arquivo_nao_do_modelo():
    """O `bruce` contou 21 aprovados e 2 cancelados num arquivo com 20 e 3.
    O resumo tem que trazer a contagem pronta, para ele não ter o que contar."""
    eventos = [ev(ANA, f"{d:02d}/09/2026", ident=f"90000{d:02d}") for d in (1, 2, 3, 4)]
    eventos += [ev(ANA, f"{d:02d}/09/2026", ident=f"90000{d:02d}", status="Cancelado") for d in (8, 9, 10)]
    r = rel(mes_inteiro(ANA), eventos)
    texto = p.resumo(r)
    assert "Eventos no arquivo: 7 (4 aprovado, 3 cancelado); 4 usados, 3 não usados" in texto


# ── o que NUNCA sai ─────────────────────────────────────────────────────────

def test_cid_e_justificativa_do_evento_nao_saem_em_lugar_nenhum():
    """Ela pediu explicitamente: não reproduzir diagnóstico. Nem na planilha
    (que é privada), nem no resumo (que vai ao contexto do agente), nem no
    histórico. Nem chegam a ser lidos."""
    r = rel([marc(ANA, 1, "09:50", "12:00", "13:00", "17:00")],
            [ev(ANA, "01/09/2026", "01/09/2026", "07:30", "10:00", motivo="Atestado Médico",
                cid="Z99.9-CID-SECRETO", justificativa="DIAGNOSTICO-SECRETO")])
    _, registros = p.identificar("e.csv", csv(CAB_EVENTOS, [ev(ANA, "01/09/2026", cid="X", justificativa="Y")]))
    assert "CID" not in registros[0] and "Justificativa" not in registros[0]
    for texto in (_todo_texto_do_xlsx(r), p.resumo(r), json.dumps(p.historico(r, "2026-10-28T00:00:00"))):
        assert "SECRETO" not in texto


def test_resumo_ao_agente_nao_tem_nome_de_pessoa():
    """O resumo vai para o contexto do modelo, que pode citá-lo em qualquer
    lugar — inclusive numa página pública, como em 21/10."""
    r = rel(mes_inteiro(ANA, ("07:55", "12:00", "13:00", "17:00"))
            + mes_inteiro(BIA, ("07:00", "12:00", "12:30", "17:30")), [ev(ANA, "02/09/2026")])
    texto = p.resumo(r)
    assert "Ana" not in texto and "Bia" not in texto and "Teste" not in texto
    assert "SETEMBRO/2026" in texto and "A2 — Atraso: 20/41 = 48,8%" in texto
    assert "Fechamento: conferido" in texto


def test_texto_do_flash_nao_vira_formula_no_excel():
    """Nome e motivo vêm do Flash. Aberto no Excel da RH, `=HYPERLINK(...)`
    viraria fórmula — o mesmo cuidado do `enviar_arquivo` (`anexos_escrita.py`)."""
    r = rel([marc(('=HYPERLINK("http://x","clique")', "000001"), 1,
                  ("07:55", "Justificada", "@SOMA(A1)", "x"), "12:00", "13:00", "17:00")])
    ws = _livro(r)["Ocorrências"]
    nome, ajuste = ws.cell(2, 2), ws.cell(2, 9)
    assert nome.value == '=HYPERLINK("http://x","clique")' and nome.data_type == "s"
    assert ajuste.value.startswith("batida ajustada: @SOMA") and ajuste.data_type == "s"


def test_caractere_de_controle_no_texto_do_flash_nao_derruba_a_planilha():
    r = rel([marc(ANA, 1, ("07:55", "Justificada", "Esque\x01cimento", "x"), "12:00", "13:00", "17:00")])
    assert _livro(r)["Ocorrências"].cell(2, 9).value == "batida ajustada: Esquecimento"


# ── fechamento ──────────────────────────────────────────────────────────────

def _r_completo():
    return rel(mes_inteiro(ANA, ("07:55", "12:00", "12:30", "17:00"), exceto=(2, 15))
               + mes_inteiro(BIA, ("07:30", "12:00", "13:00", "17:30"), exceto=(3,))
               + [marc(BIA, 3, "07:10", "17:00")],
               [ev(ANA, "02/09/2026")], ((*ANA, "(-07:42)"), (*BIA, "01:00")))


def test_fechamento_confere_e_a_conta_das_faixas_fecha():
    r = _r_completo()
    assert r.conferencias and all(c.ok for c in r.conferencias)
    for m in p.MEDIDAS:
        baixo, _, alto = p.FAIXAS[m]
        t = p.totais(r)
        neutras = sum(1 for d in r.dias if d.faixa[m] == p.NEUTRA)
        soma = sum(t[a]["identificadas"] + t[a]["justificadas"] for a in (baixo, alto)) + neutras
        assert soma == t[baixo]["denominador"]


def test_planilha_que_nao_fecha_nao_e_entregue():
    """Divergência entre a soma por pessoa e o total é defeito: a ferramenta
    recusa entregar em vez de mandar número que não fecha."""
    r = _r_completo()
    r.pessoas[0].identificadas["A2"] += 1
    assert any(not c.ok for c in p.conferir(r))
    with pytest.raises(p.ErroPonto, match="não fecha"):
        p.exigir_fechamento(r)


# ── o xlsx ──────────────────────────────────────────────────────────────────

def test_abas_e_nome_do_arquivo():
    r = _r_completo()
    assert _livro(r).sheetnames == list(p.ABAS)
    assert p.nome_do_arquivo(r) == "Controle de jornada - SETEMBRO-2026.xlsx"


def test_por_colaborador_tem_numeros_e_nenhum_rotulo_de_status():
    ws = _livro(_r_completo())["Por colaborador"]
    cab = [c.value for c in ws[1]]
    ana = dict(zip(cab, next(l for l in ws.iter_rows(min_row=2, values_only=True) if l[1] == "Ana Teste")))
    assert ana["A2 identificadas"] == 19 and isinstance(ana["A2 identificadas"], int)
    assert ana["A2 denominador"] == 19 and ana["A2 %"] == 1
    assert (ana["Dias úteis"], ana["Dias fora (evento)"], ana["Dias pendentes"]) == (21, 1, 1)
    assert ana["Saldo banco"] == "-07:42" and ana["Saldo (min)"] == -462
    assert not any(c in cab for c in ("Status", "Padrão", "Tratativa"))
    texto = _todo_texto_do_xlsx(_r_completo())
    assert "CRÍTICO" not in texto and "ATENÇÃO" not in texto


def test_ocorrencia_tem_horario_como_hora_do_dia():
    ws = _livro(_r_completo())["Ocorrências"]
    linha = next(ws.iter_rows(min_row=2))
    assert linha[4].number_format == "hh:mm" and linha[4].value == time(7, 55)
    assert linha[3].value == "A2 — Atraso"


def test_resumo_da_planilha_traz_versao_arquivos_e_quadro():
    r = _r_completo()
    ws = _livro(r)["Resumo"]
    valores = [c.value for linha in ws.iter_rows() for c in linha if c.value is not None]
    assert p.VERSAO_REGUA in valores and "m.csv" in valores
    assert all(p.rotulo(a) in valores for a in p.INDICADORES)


# ── histórico ───────────────────────────────────────────────────────────────

def test_historico_tem_so_matricula_e_numeros():
    """O JSON fica no R2 e é a base da comparação mês a mês; não precisa de
    nome, horário nem evento para isso — então não os tem."""
    r = _r_completo()
    h = p.historico(r, "2026-10-28T12:00:00+00:00")
    texto = json.dumps(h, ensure_ascii=False)
    assert "Ana" not in texto and "Bia" not in texto and "07:55" not in texto
    assert "Atestado" not in texto and "9000001" not in texto
    assert set(h["por_matricula"]) == {"1", "2"}
    assert h["por_matricula"]["1"]["A2"] == {"identificadas": 19, "justificadas": 0,
                                             "pendentes": 1, "denominador": 19}
    assert h["periodo"] == "2026-09" and h["versao_regua"] == p.VERSAO_REGUA
    assert h["totais"]["A2"]["identificadas"] == p.totais(r)["A2"]["identificadas"]


def test_caminho_do_historico_depende_da_versao_e_dos_arquivos():
    h = p.historico(_r_completo(), "x")
    c = p.caminho_historico(h)
    assert c.startswith(f"2026-09/{p.VERSAO_REGUA}-") and c.endswith(".json") and len(c.split("-")[-1]) == 13
    h2 = dict(h, arquivos=[{"nome": "outro.csv", "sha256": "0" * 64}])
    assert p.caminho_historico(h2) != c


def test_periodo_anterior_vira_o_ano():
    assert p.periodo_anterior("2026-09") == "2026-08"
    assert p.periodo_anterior("2027-01") == "2026-12"


def _hist(periodo, versao, a2, denom, gerado="2026-09-30T00:00:00"):
    totais = {a: {"identificadas": 0, "justificadas": 0, "pendentes": 0, "denominador": denom,
                  "pessoas": 0} for a in p.INDICADORES}
    totais["A2"]["identificadas"] = a2
    return {"periodo": periodo, "versao_regua": versao, "gerado_em": gerado, "totais": totais}


def test_comparacao_com_o_mes_anterior_na_mesma_regua():
    atual = _hist("2026-09", p.VERSAO_REGUA, 20, 100)
    c = p.comparar(atual, [_hist("2026-08", p.VERSAO_REGUA, 50, 100, "2026-09-01"),
                           _hist("2026-08", p.VERSAO_REGUA, 10, 100, "2026-09-05"),  # mais recente vale
                           _hist("2026-07", p.VERSAO_REGUA, 99, 100)])
    assert c.status == "comparado" and c.periodo == "2026-08"
    assert c.linhas["A2"]["identificadas"] == 10
    assert c.linhas["A2"]["variacao_pp"] == pytest.approx(10.0)


def test_comparacao_com_regua_diferente_nao_compara_e_diz_por_que():
    """Mudou a régua, mudou a constante: comparar 10% numa régua com 10% na
    outra seria dizer que nada mudou quando a régua mudou."""
    c = p.comparar(_hist("2026-09", p.VERSAO_REGUA, 20, 100), [_hist("2026-08", "2026-10-13", 50, 100)])
    assert c.status == "outra_versao" and c.linhas == {}
    assert "2026-10-13" in c.mensagem and "AGOSTO/2026" in c.mensagem and "Reprocesse" in c.mensagem


def test_sem_historico_do_mes_anterior():
    c = p.comparar(_hist("2026-09", p.VERSAO_REGUA, 20, 100), [])
    assert c.status == "sem_historico" and "AGOSTO/2026" in c.mensagem


def test_comparacao_entra_no_resumo_e_na_planilha():
    r = _r_completo()
    atual = p.historico(r, "2026-10-28")
    r.comparacao = p.comparar(atual, [_hist("2026-08", p.VERSAO_REGUA, 10, 100)])
    assert "Comparado com AGOSTO/2026" in p.resumo(r)
    ws = _livro(r)["Resumo"]
    assert "Variação (p.p.)" in [c.value for linha in ws.iter_rows() for c in linha]
